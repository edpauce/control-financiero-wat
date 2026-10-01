#!/usr/bin/env python3
"""
tools/cruzar_operaciones.py — Las dos patas de cada traspaso interno, unidas por el banco.

Lee los extractos BCP de las SEIS cuentas del grupo y aplica `lib/operaciones.py`: une por los
últimos seis dígitos del número de operación. Lo que sale es lo que el banco demuestra, no lo
que el maestro concluyó.

PARA QUÉ SIRVE, EN CONCRETO
- Un abono con glosa anónima que en realidad viene de una cuenta nuestra deja de estar «por
  identificar». Un inter-compañía que llega como «DE OTRA CUENTA» queda nombrado por este cruce.
- Inter-compañía sin depender de la categoría escrita a mano: el número de operación es
  exacto y no necesita heurísticas de fecha e importe.
- Las cuentas secundarias dejan de parecer terceros. Sus movimientos se clasificaron como
  «PAGO PROVEEDOR = gasto» porque el código no sabía que la contraparte era uno mismo.

SOLO LECTURA. No escribe maestros ni propone clasificaciones: informa, y el criterio lo pone el
conciliador con el gate delante.

TRAMPA DEL EXPORT: los .xlsx de Telecrédito declaran mal su tamaño y `openpyxl` en modo
`read_only` ve UNA fila donde hay 79 — los 18 extractos del archivo mienten igual. Se abre sin
`read_only` a propósito; si alguien lo "optimiza", el mes entero desaparece en silencio.

Uso:
  python3 tools/cruzar_operaciones.py --mes 2026-09
  python3 tools/cruzar_operaciones.py --todo --salida informe.md
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl                                        # noqa: E402
from lib import empresas as EMP                        # noqa: E402
from lib import operaciones as OP                      # noqa: E402

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
# El mismo export sale en tres envases: .xlsx (11 columnas), .csv (7, coma) y .txt (7, pipe).
# Mayo salió en los dos últimos, así que mirar solo .xlsx dejaba el mes entero fuera y el
# recuento de pares parecía un defecto del cruce.
PATRONES = ("0[12]. */01. EECC/**/*.xlsx", "0[12]. */01. EECC/**/*.csv",
            "0[12]. */01. EECC/**/*.txt")


def _fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})", str(v or ""))
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


def _indices(cabecera) -> dict | None:
    """Posición de las 4 columnas que importan, buscadas POR NOMBRE.

    El .csv trae 7 columnas y el .xlsx 11, así que la operación está en la 5 o en la 6 según el
    envase. Fijar el índice funcionaba hasta que alguien exportaba en el otro formato — y fallaba
    en silencio leyendo la columna equivocada, que es peor que no leer. Además la codificación
    latin-1 destroza los acentos («Descripci?n»), por eso se compara por prefijo.
    """
    ix = {}
    for i, c in enumerate(str(x or "").strip().lower() for x in cabecera):
        if c == "fecha":
            ix.setdefault("fecha", i)
        elif c.startswith("descripci"):
            ix.setdefault("desc", i)
        elif c == "monto":
            ix.setdefault("monto", i)
        elif c.startswith("operaci") and "hora" not in c:
            ix.setdefault("op", i)
    return ix if len(ix) == 4 else None


def _monto(v):
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    txt = str(v or "").replace(",", "").strip()
    try:
        return round(float(txt), 2)
    except ValueError:
        return None


def _filas(ruta: str):
    """`(cabecera_cuenta, filas)` de cualquiera de los tres envases, o `(None, [])`."""
    if ruta.lower().endswith(".xlsx"):
        try:
            # sin read_only A PROPÓSITO: estos exports declaran mal su tamaño (ver docstring)
            ws = openpyxl.load_workbook(ruta, data_only=True).worksheets[0]
        except Exception:
            return None, []
        return ws["B1"].value, [list(f) for f in ws.iter_rows(values_only=True)]
    import csv
    try:
        with open(ruta, encoding="latin-1", newline="") as fh:
            crudo = fh.read()
    except Exception:
        return None, []
    sep = "|" if crudo.count("|") > crudo.count(",") else ","
    filas = list(csv.reader(crudo.splitlines(), delimiter=sep))
    if not filas or len(filas[0]) < 2:
        return None, []
    return filas[0][1], filas


def leer_extracto(ruta: str) -> list[dict]:
    """Movimientos de un export de Telecrédito, o [] si el archivo no tiene esa forma."""
    cabecera_cuenta, filas = _filas(ruta)
    cta = EMP.cuenta_del_grupo(cabecera_cuenta)
    if not cta or not filas:
        return []                                       # no es una cuenta del grupo
    emp, cuenta, moneda, _ = cta
    ix = fila_cab = None
    for n, f in enumerate(filas[:12]):                  # la cabecera real no siempre va en la 5
        ix = _indices(f)
        if ix:
            fila_cab = n
            break
    if not ix:
        return []                                       # formato desconocido: no adivinar columnas
    out = []
    for f in filas[fila_cab + 1:]:
        if len(f) <= max(ix.values()):
            continue
        fe, mo = _fecha(f[ix["fecha"]]), _monto(f[ix["monto"]])
        if fe is None or mo is None:
            continue
        out.append({"cuenta": cuenta, "empresa": emp, "moneda": moneda, "fecha": fe,
                    "monto": mo, "glosa": str(f[ix["desc"]] or "")[:30],
                    "op": str(f[ix["op"]] or ""), "archivo": os.path.basename(ruta)})
    return out


def movimientos(mes: str | None = None) -> tuple[list[dict], list[str]]:
    """Todos los movimientos del grupo, DEDUPLICADOS por (cuenta, fecha, monto, op).

    El archivo guarda quincenas parciales y el mes completo del mismo periodo. Deduplicar por
    identidad del movimiento —y no por archivo— hace que solapar exports no cuente doble.
    """
    vistos, out, leidos = {}, [], []
    rutas = sorted({r for pat in PATRONES for r in glob.glob(os.path.join(RAIZ, pat), recursive=True)})
    for ruta in rutas:
        movs = leer_extracto(ruta)
        if not movs:
            continue
        leidos.append(os.path.basename(ruta))
        for m in movs:
            if mes and m["fecha"].strftime("%Y-%m") != mes:
                continue
            vistos.setdefault((m["cuenta"], m["fecha"], m["monto"], m["op"]), m)
    out = list(vistos.values())
    out.sort(key=lambda m: (m["fecha"], m["cuenta"]))
    return out, leidos


def informe(pares, descartados, movs, mes) -> str:
    L = [f"# Traspasos internos por número de operación — {mes or 'todo el archivo'}", "",
         f"{len(movs)} movimientos de {len({m['cuenta'] for m in movs})} cuentas del grupo. "
         f"Unidos por los **últimos 6 dígitos** de la operación: los dos primeros son el canal "
         f"(`01…` carga, `02…` abona, `00…` en la constancia), no el número.", ""]
    for tipo, titulo in ((OP.INTERCIA, "Inter-compañía"), (OP.PROPIO, "Entre cuentas propias")):
        grupo = [p for p in pares if p["tipo"] == tipo]
        if not grupo:
            continue
        L += [f"## {titulo}: {len(grupo)}", "", "| Operación | Fecha | Patas | Estado |", "|---|---|---|---|"]
        for p in grupo:
            patas = " · ".join(f"{q['cuenta'][-2:]} {q['monto']:+,.2f} {q['moneda']}"
                               for q in p["patas"])
            # con `dict` las tres ramas se evalúan siempre y el f-string del neto revienta
            # en los pares de moneda mixta, donde `neto` es None a propósito
            if p["estado"] == OP.NEUTRO:
                estado = "✅ neutro"
            elif p["estado"] == OP.CAMBIO:
                estado = f"💱 cambio a {p['tasa']}" if p["tasa"] else "💱 cambio de moneda"
            else:
                estado = f"⚠️ neto {p['neto']:+,.2f}"
            L.append(f"| …{p['suf']} | {p['patas'][0]['fecha']:%d/%m} | {patas} | {estado} |")
        L.append("")
    anon = [p for p in pares for q in p["patas"]
            if q["monto"] > 0 and ("OTRA CUENTA" in q["glosa"].upper() or q["glosa"].startswith("TRAN"))]
    if anon:
        L += [f"## Glosas anónimas que este cruce resolvió: {len(anon)}", "",
              "Abonos que sin el número de operación quedarían «por identificar» en la pila del "
              "cobrador, y que son movimiento nuestro.", ""]
        for p in anon:
            q = next(x for x in p["patas"] if x["monto"] > 0)
            L.append(f"- {q['fecha']:%d/%m} · {q['cuenta']} · {q['monto']:+,.2f} · "
                     f"«{q['glosa']}» → op …{p['suf']} ({p['tipo']})")
        L.append("")
    if descartados:
        L += [f"## Descartados: {len(descartados)}", "",
              "Sufijos que coinciden pero NO son un par. Se listan en vez de tirarse: un "
              "descarte invisible es un hueco que nadie audita.", ""]
        L += [f"- …{d['suf']} — {d['motivo']}" for d in descartados]
    return "\n".join(L) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--mes", help="AAAA-MM; por defecto, todo el archivo")
    ap.add_argument("--todo", action="store_true", help="explícito: sin filtro de mes")
    ap.add_argument("--salida", help="ruta del informe .md (si no, va a pantalla)")
    a = ap.parse_args()

    movs, leidos = movimientos(None if a.todo else a.mes)
    if not movs:
        print("⚠️  Ningún movimiento: revisa el mes o que los extractos estén archivados.")
        return 1
    pares, descartados = OP.emparejar(movs)
    txt = informe(pares, descartados, movs, None if a.todo else a.mes)

    if a.salida:
        with open(a.salida, "w", encoding="utf-8") as fh:
            fh.write(txt)
        print(f"💾 {a.salida}")
    else:
        print(txt)
    mal = [p for p in pares if p["estado"] == OP.DESCUADRE]
    print(f"   {len(leidos)} extractos · {len(movs)} movimientos · {len(pares)} pares "
          f"({sum(p['tipo'] == OP.INTERCIA for p in pares)} inter-cía) · {len(mal)} descuadres")
    return 1 if mal else 0


if __name__ == "__main__":
    raise SystemExit(main())
