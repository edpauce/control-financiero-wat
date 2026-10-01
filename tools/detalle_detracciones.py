#!/usr/bin/env python3
"""
tools/detalle_detracciones.py — Detalle SPOT del MES factura por factura.

Aplica la regla de negocio (lib/empresas.py):
  - Cliente pagó el TOTAL  → AUTO-DETRACCIÓN nuestra (obligación) → a 09 Pendientes.
  - Cliente pagó NETO      → el CLIENTE deposita el 12% (verificar BN; puede caer el mes siguiente).
  - Total ≤ umbral / exon. → Sin Detracción.
Cruza: emitidas efact × estado de cobro (COBROS_{mes}) × depósito (SUNAT del mes).

Escribe las hojas 'Detalle {Mes} ALT/MER' en el BASE SEGUIMIENTO_DETRACCIONES_{anio}.xlsx (con
backup .bak; persisten al regenerar) + reporte MD. Correr DESPUÉS de integrar_cobros y ANTES de
integrar_detracciones.

Uso:  python3 tools/detalle_detracciones.py --mes 2026-07
"""
from __future__ import annotations
import csv
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import cierre as CI     # noqa: E402
from lib import empresas as EMPX  # noqa: E402
from lib import constancias as CONST  # noqa: E402

ROOT = CI.ROOT
COLS = ["Periodo", "Fecha Emisión", "Comprobante", "Cliente", "Total CP S/", "Detr Esperada (12%)",
        "Detr Depositada (BN/SUNAT)", "N° Constancia (F. Dep.)", "Forma de Pago", "Estado SPOT",
        "Observación"]
HOJA_COB = {"ALTAMIRA": "ALTAMIRA - Cobros", "MERAK": "MERAK - Cobros"}
FILL = {"✅": "E2EFDA", "🔴": "F8CBAD", "🟡": "FFF2CC", "🟠": "FDE9D9", "Sin": "F2F2F2",
        "Anu": "D9D9D9", "Por": "DDEBF7", "⚠️": "FCE4D6"}


def emitidas(cfg, empresa):
    f = CI.first(CI.rutas(cfg)["efact"][empresa])
    if not f:
        return []
    ws = openpyxl.load_workbook(f, data_only=True).active
    hdr = [c.value for c in ws[1]]
    C = lambda n: next(i for i, h in enumerate(hdr) if n.lower() in str(h).lower())
    iN, iC, iF, iT, iTp = C("Comprobante"), C("Razón"), C("Emisi"), C("Total en Soles"), C("Tipo Doc")
    iR = next((i for i, h in enumerate(hdr) if str(h or "").strip().upper() == "RUC"), None)
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[iN]:
            continue
        out.append({"comp": CI.comp_norm(r[iN]), "cli": str(r[iC]).split(" - ")[0][:34], "f": r[iF],
                    "tot": CI.num(r[iT]), "tp": str(r[iTp]),
                    "ruc": str(r[iR]).strip() if iR is not None and r[iR] else ""})
    return out


def constancias_sunat(cfg, empresa):
    """
    P3a: las constancias del CSV SUNAT con TODAS sus claves de match, no solo el monto.
    Devuelve (por_comp, todas):
      por_comp: comprobante_norm → [constancia…]  (join exacto, el camino feliz)
      todas:    lista completa — permite el join INVERSO: una constancia cuyo comprobante no
                casa con ninguna factura (un cliente consignó E001-155 en vez de E001-1552) ya
                no desaparece en silencio; sale como huérfana con RUC+monto+fecha, la materia
                prima del fallback de P3b.
    Cada constancia: {constancia, comp, ruc_adq, adquiriente, fecha, monto, periodo}.
    """
    full = CI.first(CI.rutas(cfg)["sunat"][empresa])
    if not full:
        return {}, []
    todas = _parse_constancias_csv(full)
    por_comp = {}
    for c in todas:
        por_comp.setdefault(c["comp"], []).append(c)
    return por_comp, todas


def _parse_constancias_csv(full):
    """Parsea UN csv de constancias SUNAT → lista de constancias (claves de constancias_sunat)."""
    raw = open(full, "rb").read()
    for e in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            t = raw.decode(e); break
        except UnicodeDecodeError:
            continue
    rows = list(csv.reader(t.splitlines()))
    if not rows:
        return []
    hdr = rows[0]

    def C(n):
        return next((i for i, h in enumerate(hdr) if n.lower() in h.lower()), None)
    iS, iN, iM = C("Serie de Comprobante"), C("Numero de Comprobante"), C("Monto Deposito")
    # 'Numero de Documento Adquiriente' explícito: el CSV también trae 'TIPO de Documento
    # Adquiriente' (valor '06'), y un substring genérico casa primero con esa.
    iK, iR, iA = (C("Numero Constancia"), C("Numero de Documento Adquiriente"),
                  C("Razon Social del Adquiriente"))
    iF, iP = C("Fecha Pago"), C("Periodo Tributario")
    if iN is None or iM is None:
        return []

    def g(r, i):
        return r[i].strip() if i is not None and i < len(r) else ""
    todas = []
    for r in rows[1:]:
        if not any(c.strip() for c in r):
            continue
        serie = g(r, iS) or "F001"
        todas.append({"constancia": g(r, iK), "comp": CI.comp_norm(f"{serie}-{g(r, iN)}"),
                      "ruc_adq": g(r, iR), "adquiriente": g(r, iA)[:34],
                      "fecha": g(r, iF), "monto": CI.num(g(r, iM)), "periodo": g(r, iP)})
    return todas


def candado_hoja09(wb, cfg):
    """
    Candado de la hoja '09 Pendientes Auto-Det': una fila PENDIENTE cuyo comprobante YA tiene
    constancia SUNAT (en CUALQUIER mes del año) es un falso pendiente. Así se coló una factura
    que el cliente ya había depositado un mes antes: la lista manual
    no cruzó contra las constancias. Avisa — no cambia estados solo (decisión del usuario).
    """
    import glob as _g
    if "09 Pendientes Auto-Det" not in wb.sheetnames:
        return []
    por_emp = {}
    for emp in ("ALTAMIRA", "MERAK"):
        d, ab = CI.DIR_EMP[emp], CI.ABREV_EMP[emp]
        m = {}
        for f in sorted(_g.glob(os.path.join(CI.ROOT, d, "03. DETRACCIONES", str(cfg.ANIO),
                                             "02. Constancias_SUNAT", f"DETR-SUNAT_{ab}_*.csv"))):
            for c in _parse_constancias_csv(f):
                if c["comp"]:
                    m[c["comp"]] = c
        por_emp[emp] = m
    avisos = []
    ws = wb["09 Pendientes Auto-Det"]
    for r in range(2, ws.max_row + 1):
        emp = str(ws.cell(r, 1).value or "").strip().upper()
        comp = CI.comp_norm(ws.cell(r, 5).value)
        estado = str(ws.cell(r, 8).value or "").upper()
        if emp not in ("ALTAMIRA", "MERAK") or not comp or not estado.startswith("PENDIENTE"):
            continue
        c = por_emp[emp].get(comp)
        if c:
            avisos.append(f"{emp} {comp}: figura PENDIENTE en hoja 09 pero YA tiene constancia "
                          f"{c['constancia']} ({c['fecha']}, S/{c['monto']:,.2f}) — márcala DEPOSITADO")
    for a in avisos:
        print("  ⚠️", a)
    return avisos


def sunat_declaradas(cfg, empresa):
    """comprobante → monto depositado (compatibilidad; deriva de constancias_sunat)."""
    por_comp, _ = constancias_sunat(cfg, empresa)
    return {k: sum(c["monto"] for c in v) for k, v in por_comp.items()}


def bn_abonos(cfg, empresa):
    """
    Abonos de la cuenta BN de detracciones (NOT.ABON): [{fecha, ruc, doc, monto}].
    Es el HECHO (el dinero entró); la constancia es el REGISTRO. El RUC del depositante viene
    en la columna 5 (a veces vacío). Vacío si el mes no tiene BN csv (mayo cerró sin él).
    """
    f = CI.first(CI.rutas(cfg)["bn"][empresa])
    if not f or not f.lower().endswith((".csv", ".txt")):
        return []
    t = open(f, "rb").read().decode("latin-1", "replace")
    rows = list(csv.reader(t.splitlines()))
    hj = next((i for i, r in enumerate(rows) if r and str(r[0]).strip() == "Nro"), None)
    if hj is None:
        return []
    out = []
    for r in rows[hj + 1:]:
        r = (list(r) + [""] * 8)[:8]
        if "ABON" not in str(r[2]).upper() or not CI.num(r[7]):
            continue
        out.append({"fecha": str(r[1]).strip(), "ruc": str(r[4]).strip(),
                    "doc": str(r[3]).strip(), "monto": CI.num(r[7])})
    return out


def _bn_confirma(constancia, abonos):
    """¿El BN tiene un abono que confirme esta constancia? RUC (si ambos lo traen) + monto exacto."""
    for a in abonos:
        if abs(a["monto"] - constancia["monto"]) > 0.01:
            continue
        if constancia["ruc_adq"] and a["ruc"] and constancia["ruc_adq"] != a["ruc"]:
            continue
        return a
    return None


# ─── P3b: fallback RUC + monto + fecha, con desambiguación ───────────────────
# Tolerancia de monto ±1.00 y NO el ±0.01 del plan: SUNAT deposita la detracción REDONDEADA AL
# SOL ENTERO (esperada 1,598.79 → depositó 1,599.00; esperada
# 1,998.28 → depositó 1,998.00). Con ±0.01 el fallback no casaría ningún caso real.
TOL_MONTO = 1.00
TOL_DIAS = 3


def _dias(a, b) -> int | None:
    from lib import sources as SRC
    da, db = SRC.parse_date(a), SRC.parse_date(b)
    if not da or not db:
        return None
    return abs((da - db).days)


def fallback_p3b(filas, meta, huerfanas):
    """
    Intenta casar cada constancia huérfana con una factura sin constancia por
    RUC + monto (±1.00, redondeo SUNAT) + fecha de COBRO ±3 días (la detracción se deposita
    cerca del pago, no de la emisión; sin fecha de cobro NO hay match — conservador).

    REGLA DURA (plan CEO + revisor adversarial): esto NUNCA auto-cierra. Un candidato único
    marca la factura 🟠 'Constancia por confirmar' con la propuesta visible; múltiples
    candidatos en cualquier dirección marcan ⚠️ AMBIGUO y no se propone nada. La columna
    'Detr Depositada' NO se toca: cerrar sobre un match difuso es peor que no cerrar.
    """
    pares = []                                     # (idx_huerfana, idx_fila) válidos
    for hi, h in enumerate(huerfanas):
        for fi, m in enumerate(meta):
            if m is None or m["dep"] > 0 or m["esp"] <= 0:
                continue
            if not h["ruc_adq"] or h["ruc_adq"] != m["ruc"]:
                continue
            if abs(h["monto"] - m["esp"]) > TOL_MONTO:
                continue
            d = _dias(h["fecha"], m["fcobro"])
            if d is None or d > TOL_DIAS:
                continue
            pares.append((hi, fi))

    from collections import Counter
    por_h = Counter(hi for hi, _ in pares)
    por_f = Counter(fi for _, fi in pares)
    for hi, fi in pares:
        h, fila = huerfanas[hi], filas[fi]
        if por_h[hi] == 1 and por_f[fi] == 1:      # único en AMBAS direcciones
            fila[7] = f"¿{h['constancia']} ({h['fecha']})?"
            fila[9] = "🟠 Constancia por confirmar"
            fila[10] += (f" · constancia {h['constancia']} consignada como {h['comp']} — "
                         f"RUC+monto+fecha cuadran; confirmar con el cliente, NO auto-cerrada")
            h["p3b"] = f"candidata única de {fila[2]} (comprobante mal consignado)"
        else:
            fila[10] += (f" · ⚠️ {por_f[fi]} constancia(s) candidatas — AMBIGUO, resolver a mano")
            h["p3b"] = f"AMBIGUA ({por_h[hi]} factura(s) candidatas) — resolver a mano"
    return filas, huerfanas


def cobros(cfg, empresa):
    path = os.path.join(ROOT, "04. FINANZAS", "08. COBROS", f"COBROS_{cfg.PERIODO}.xlsx")
    if not os.path.exists(path):
        sys.exit(f"❌ Falta COBROS_{cfg.PERIODO}.xlsx — correr integrar_cobros primero.")
    ws = openpyxl.load_workbook(path, data_only=True)[HOJA_COB[empresa]]
    d = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        c = str(r[2] or "")
        if c.startswith(("F001", "E001")):
            d[CI.comp_norm(c)] = {"estado": str(r[12] or ""), "cobrado": CI.num(r[10]),
                                  "total": CI.num(r[4]), "fcobro": r[9]}
    return d


def clasificar(total, esp, cob):
    estado = cob.get("estado", "") if cob else ""
    cobrado = cob.get("cobrado", 0) if cob else 0
    if esp == 0:
        return "—", "Sin Detracción", "operación no sujeta / ≤ umbral"
    if estado.lower() == "anulado":
        return "—", "Anulado", "factura anulada (NC)"
    if estado.lower() != "cobrado":
        return "—", "Por cobrar", f"aún no cobrada (estado {estado or '—'}) — sin obligación todavía"
    if abs(cobrado - total) < 2:
        return "completo", "AUTO", "pagó el total → auto-detracción nuestra"
    if abs(cobrado - (total - esp)) < 2:
        return "neto", "CLIENTE", "pagó neto → el cliente deposita el 12%"
    return "parcial", "REVISAR", f"cobró {cobrado:,.0f} (≠ total {total:,.0f} ni neto {total - esp:,.0f})"


def build(cfg, empresa):
    """(filas, huerfanas): filas del detalle + constancias que no casaron con ninguna factura."""
    por_comp, todas = constancias_sunat(cfg, empresa)
    abonos = bn_abonos(cfg, empresa)
    cob = cobros(cfg, empresa)
    per = f"{cfg.ANIO}{cfg.MES:02d}"
    filas, usadas, meta = [], set(), []
    # Hoja 09: facturas cobradas BRUTO (auto-detracción NUESTRA). Es la única marca de autoría que
    # sobrevive al cierre por constancia (ver abajo).
    bruto = CONST.cobradas_bruto()
    for x in sorted(emitidas(cfg, empresa), key=lambda y: str(y["comp"])):
        if x["tp"] != "Factura":
            filas.append([per, x["f"], x["comp"], x["cli"], x["tot"], 0, 0, "", "—", "🔘 NC",
                          "nota de crédito"])
            meta.append(None)                      # las NC no participan del fallback
            continue
        esp = EMPX.detr_esperada(x["tot"])
        consts = por_comp.get(x["comp"], [])
        for c in consts:
            usadas.add(id(c))
        dep = sum(c["monto"] for c in consts)
        # P3a: la constancia deja de ser un monto anónimo — nº y fecha van a su columna,
        # y el BN confirma (o no) que el dinero realmente entró.
        ref = " · ".join(f"{c['constancia']} ({c['fecha']})" for c in consts if c["constancia"])
        forma, via, obs = clasificar(x["tot"], esp, cob.get(x["comp"]))
        if via == "AUTO" and dep > 0 and x["comp"] not in bruto:
            # COBROS cierra por constancia con Monto Cobrado = TOTAL (neto en BCP + 12% en
            # BN), así que 'cobrado == total' ya NO prueba pago BRUTO. Si la hoja 09 no la tiene
            # marcada BRUTO, la constancia es depósito del CLIENTE (una factura pagada NETO
            # salía como 'auto-detracción nuestra'). La autoría real se prueba con la salida PAGO DETRAC del BCP.
            forma, via = "neto", "CLIENTE"
            obs = "cerrada por constancia → cobro NETO: el 12% lo depositó el cliente (sin marca BRUTO en hoja 09)"
        if consts and abonos:
            conf = [c for c in consts if _bn_confirma(c, abonos)]
            if conf:
                obs += " · abono BN ✓"
        if via == "Sin Detracción":
            est = "🔘 Sin Detracción"
        elif via == "Anulado":
            est = "🔘 Anulado"
        elif via == "Por cobrar":
            est = "🔵 Por cobrar"
        elif via == "AUTO":
            est = "✅ Auto-detraído" if dep > 0 else "🔴 Auto-detraer"
            obs += (" · 12% ya en BN" if dep > 0 else f" · pendiente S/{esp:,.0f} (→ 09 Pendientes)")
        elif via == "CLIENTE":
            est = "✅ Cliente depositó" if dep > 0 else "🟡 Cliente deposita"
            obs += (" · confirmado en BN/SUNAT" if dep > 0 else " · verificar BN (puede caer el mes sgte.)")
        else:
            est = "⚠️ Revisar"
        filas.append([per, x["f"], x["comp"], x["cli"], x["tot"], esp, dep, ref, forma, est, obs])
        c_x = cob.get(x["comp"]) or {}
        meta.append({"ruc": x["ruc"], "esp": esp, "dep": dep, "fcobro": c_x.get("fcobro")})

    # Join inverso: constancias del CSV que ninguna factura reclamó. Puede ser una regularización
    # de un mes anterior (legítimo: el periodo tributario lo dice) o un comprobante mal consignado
    # por el cliente. P3a las hace VISIBLES; P3b propone (nunca cierra) por fallback.
    huerfanas = [c for c in todas if id(c) not in usadas]
    filas, huerfanas = fallback_p3b(filas, meta, huerfanas)
    return filas, huerfanas


def escribir_hoja(wb, cfg, empresa, filas):
    nombre = f"Detalle {cfg.MES_NOMBRE} {CI.ABREV_EMP[empresa]}"
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)
    ws.append([f"{empresa} — Detalle Detracciones {cfg.MES_NOMBRE.upper()} {cfg.ANIO} "
               f"(SPOT factura×cobro×BN/SUNAT)"])
    ws.append(["Regla: pagó total → auto-detracción nuestra · pagó neto → el cliente deposita (verificar BN)"])
    ws.append([])
    ws.append(COLS)
    HDR = PatternFill("solid", fgColor="1F4E79")
    for c in range(1, len(COLS) + 1):
        cell = ws.cell(4, c); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    tv = te = td = 0.0
    r0 = 5
    for i, f in enumerate(filas):
        ws.append(f)
        est = str(f[9])
        key = next((k for k in FILL if est.startswith(k) or k in est), None)
        if key:
            fill = PatternFill("solid", fgColor=FILL[key])
            for c in range(1, len(COLS) + 1):
                ws.cell(r0 + i, c).fill = fill
        for c in (5, 6, 7):
            ws.cell(r0 + i, c).number_format = '#,##0.00'
        if not est.endswith("NC"):
            tv += f[4]; te += f[5]; td += f[6]
    ws.append(["", "", "", "TOTAL", round(tv, 2), round(te, 2), round(td, 2), "", "", "", ""])
    for c in range(1, len(COLS) + 1):
        cell = ws.cell(ws.max_row, c); cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="B4C7E7")
        if c in (5, 6, 7):
            cell.number_format = '#,##0.00'
    for c, w in zip("ABCDEFGHIJK", (8, 12, 12, 30, 12, 13, 15, 22, 11, 18, 44)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A5"
    return tv, te, td, filas


def main():
    cfg = CI.cargar_config(CI.arg_mes())
    base = os.path.join(ROOT, "04. FINANZAS", "10. DETRACCIONES", f"SEGUIMIENTO_DETRACCIONES_{cfg.ANIO}.xlsx")
    if not os.path.exists(base):
        sys.exit(f"❌ Falta el master base {os.path.basename(base)}.")
    # Respaldo antes de escribir in-place: un slot por periodo, no se pisa si ya existe.
    CI.respaldar(base, cfg.PERIODO)
    wb = openpyxl.load_workbook(base)
    rep = [f"# Detalle de Detracciones — {cfg.MES_NOMBRE.upper()} {cfg.ANIO}", "",
           "SPOT factura por factura aplicando la regla: **pagó total → auto-detracción nuestra**; "
           "**pagó neto → el cliente deposita** (verificar BN, puede caer el mes siguiente).", ""]
    for empresa in ("ALTAMIRA", "MERAK"):
        filas_mes, huerfanas = build(cfg, empresa)
        tv, te, td, filas = escribir_hoja(wb, cfg, empresa, filas_mes)
        auto = [f for f in filas if "Auto-detraer" in f[9]]
        cli = [f for f in filas if "Cliente deposita" in f[9]]
        rep += [f"## {empresa} — {len([f for f in filas if not str(f[9]).endswith('NC')])} facturas · "
                f"ventas S/{tv:,.2f}",
                f"- Detr. esperada S/{te:,.2f} · depositada (BN/SUNAT) **S/{td:,.2f}**",
                f"- 🔴 Auto-detracción NUESTRA pendiente: **{len(auto)}** (S/{sum(f[5] for f in auto):,.2f})",
                f"- 🟡 Cliente deposita (verificar BN): **{len(cli)}** (S/{sum(f[5] for f in cli):,.2f})", ""]
        for lbl, grp in (("🔴 Auto-detraer (nosotros)", auto), ("🟡 Cliente deposita (verificar)", cli)):
            if grp:
                rep += [f"**{lbl}:**", "| Comprobante | Cliente | Total | 12% |", "|---|---|---|---|"]
                rep += [f"| {f[2]} | {f[3][:26]} | S/{f[4]:,.0f} | S/{f[5]:,.0f} |" for f in grp]
                rep.append("")
        if huerfanas:
            rep += [f"**⚠️ Constancias sin factura del mes que case ({len(huerfanas)})** — "
                    f"regularización de un mes previo o comprobante mal consignado por el cliente:",
                    "| Constancia | Comp. consignado | Adquiriente | RUC | Monto | F. Pago | Periodo trib. |",
                    "|---|---|---|---|---|---|---|"]
            rep += [f"| {h['constancia']} | {h['comp']} | {h['adquiriente'][:24]} | {h['ruc_adq']} "
                    f"| S/{h['monto']:,.2f} | {h['fecha']} | {h['periodo']}"
                    f"{' → ' + h['p3b'] if h.get('p3b') else ''} |" for h in huerfanas]
            rep.append("")
        print(f"  {empresa}: {len(filas)} facturas · esperada S/{te:,.0f} · depositada S/{td:,.0f} · "
              f"auto-detraer {len(auto)} · cliente-deposita {len(cli)}")
        for h in huerfanas:
            extra = f" → {h['p3b']}" if h.get("p3b") else ""
            print(f"      ⚠️ constancia huérfana {h['constancia']}: comp {h['comp']} "
                  f"S/{h['monto']:,.2f} {h['fecha']} · {h['adquiriente'][:26]} "
                  f"(periodo {h['periodo']}){extra}")
    avisos09 = candado_hoja09(wb, cfg)
    if avisos09:
        rep += ["## ⚠️ Candado hoja 09 — pendientes que YA tienen constancia", ""]
        rep += [f"- {a}" for a in avisos09] + [""]
    wb.save(base)
    print(f"\n✅ Hojas 'Detalle {cfg.MES_NOMBRE} ALT/MER' escritas en {os.path.basename(base)} (base)")
    out = os.path.join(ROOT, "04. FINANZAS", "05. INFORMES",
                       f"DETALLE_DETRACCIONES_{cfg.MES_NOMBRE.upper()}_{cfg.ANIO}.md")
    open(out, "w", encoding="utf-8").write("\n".join(rep))
    print(f"✅ Reporte: {os.path.basename(out)}")


if __name__ == "__main__":
    main()
