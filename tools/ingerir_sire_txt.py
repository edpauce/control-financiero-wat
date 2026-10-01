#!/usr/bin/env python3
"""
tools/ingerir_sire_txt.py — Ingesta un TXT del SIRE (RVIE ventas / RCE compras) del contador,
lo estructura en un xlsx y lo CRUZA contra nuestros maestros:

  · RCE (compras) → cruza contra PROVEEDORES (CxP): qué declaró el contador que NO capturamos
    de Mail (aseguradoras, Makro, etc. no mandan XML) — la fuente oficial del crédito fiscal.
  · RVIE (ventas) → cruza contra COBROS: confirma que lo declarado = lo que tenemos, y caza
    ventas que se nos escapan (E001 por SUNAT SOL).

El TXT es la fuente OFICIAL (lo declarado). No toca los maestros; produce el xlsx estructurado
y un reporte de cruce. Que el pagador/cobrador decidan qué incorporar.

Uso:  python3 tools/ingerir_sire_txt.py "ruta/al/RCE_...txt" [--empresa ALT]
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl                              # noqa: E402
from lib import cierre as CI                 # noqa: E402
from lib import sire_txt as ST               # noqa: E402
from lib import policy as POL
from lib import formato as FMT               # noqa: E402
from lib import sources as S                 # noqa: E402

HOJA_COB = {"ALT": "ALTAMIRA - Cobros", "MER": "MERAK - Cobros"}


def _comps_maestro(ruta, hojas) -> set:
    """Comprobantes ya presentes en un maestro (para el cruce), normalizados F001-1797."""
    import re
    out = set()
    if not ruta or not os.path.exists(ruta):
        return out
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    for h in hojas:
        if h not in wb.sheetnames:
            continue
        for r in wb[h].iter_rows(values_only=True):
            for v in (r or []):
                m = re.match(r"([A-Z]\d{3,4})-0*(\d+)", str(v or ""))
                if m:
                    out.add(f"{m.group(1)}-{m.group(2)}")
    return out


def _norm(comp: str) -> str:
    import re
    m = re.match(r"([A-Za-z]+\d*)-0*(\d+)", comp)
    return f"{m.group(1)}-{m.group(2)}" if m else comp


def escribir_xlsx(libro, regs, periodo, emp):
    ws_tit = "Compras (RCE)" if libro == "RCE" else "Ventas (RVIE)"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = ws_tit
    ws.append([f"{libro} {periodo} · fuente oficial SUNAT (declarado por el contador)"])
    cab = ["Comprobante", "Tipo", "Fecha", "RUC contraparte", "Contraparte",
           "BI Gravada S/", "IGV S/", "Total S/", "Moneda", "Anula/Modifica"]
    ws.append(cab)
    FMT.encabezado(ws, 2, len(cab))
    for r in regs:
        ws.append([r["comp"], "NC" if r["es_nc"] else "Factura", r["fecha"], r["ruc"],
                   r["contraparte"][:40], r["bi"], r["igv"], r["total"], r["moneda"],
                   r.get("modifica", "")])
    t = ST.totales(regs)
    fila = ws.max_row + 1
    ws.append([f"TOTAL ({t['n']})", "", "", "", "", t["bi"], t["igv"], t["total"], "", ""])
    from openpyxl.styles import Font, PatternFill
    for c in range(1, len(cab) + 1):
        ws.cell(fila, c).font = Font(bold=True)
        ws.cell(fila, c).fill = PatternFill("solid", fgColor=FMT.TOTAL_FILL)
    FMT.pulir_hoja(ws)
    carpeta = os.path.join(S.ROOT, "01. ALTAMIRA" if emp == "ALT" else "02. MERAK",
                           "05. TRIBUTARIO_SIRE" if emp == "ALT" else "05. TRIBUTARIO_SIRE",
                           "2026")
    os.makedirs(carpeta, exist_ok=True)
    out = os.path.join(carpeta, f"SIRE_{libro}_{emp}_{periodo}.xlsx")
    if not POL.exigir(POL.evaluar(POL.Accion(
            tool="ingerir_sire_txt", maestro=out, empresa=None, capacidad="cobros.escribir",
            periodo=POL.periodo_de(out), aplicar=True,
            monto=max((abs(CI.num(r["total"])) for r in regs), default=0.0),
            total=t["total"],
            evidencia="TXT SIRE entregado por el contador",
            forzar="--forzar" in sys.argv))):
        sys.exit("(no se escribió)")
    wb.save(out)
    return out, t


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    ruta = args[0]
    emp = sys.argv[sys.argv.index("--empresa") + 1].upper() if "--empresa" in sys.argv else "ALT"
    libro, regs = ST.parse(ruta)
    if libro == "DESCONOCIDO":
        sys.exit("❌ No reconozco el TXT como RVIE ni RCE (¿encabezado distinto?).")
    # periodo AAAA-MM desde la primera fecha DD/MM/AAAA legible (algunos comprobantes no la traen)
    import re as _re
    periodo = "sin-periodo"
    for r in regs:
        m = _re.match(r"(\d{2})/(\d{2})/(\d{4})", r["fecha"])
        if m:
            periodo = f"{m.group(3)}-{m.group(2)}"
            break
    out, t = escribir_xlsx(libro, regs, periodo, emp)
    print(f"📥 {libro} {periodo} ({emp}): {t['n']} comprobantes · IGV S/{t['igv']:,.2f} · "
          f"total S/{t['total']:,.2f}")
    print(f"   xlsx: {os.path.basename(out)}")

    # ── cruce contra nuestros maestros ──
    if libro == "RCE":
        prov_master = S.latest("proveedores")
        nuestros = _comps_maestro(prov_master, ["ALTAMIRA", "MERAK", "Compras", "Proveedores"]) \
            if prov_master else set()
        faltan = [r for r in regs if _norm(r["comp"]) not in nuestros]
        print(f"\n🔎 Cruce RCE vs CxP (PROVEEDORES): {len(regs) - len(faltan)} ya capturados · "
              f"{len(faltan)} que el contador declaró y NO teníamos:")
        for r in faltan:
            print(f"   + {r['comp']:14} {r['contraparte'][:34]:34} IGV S/{r['igv']:>7,.2f} "
                  f"Total S/{r['total']:>8,.2f}")
    else:
        cob = S.latest("cobros")
        nuestros = _comps_maestro(cob, list(HOJA_COB.values())) if cob else set()
        faltan = [r for r in regs if _norm(r["comp"]) not in nuestros and not r["es_nc"]]
        print(f"\n🔎 Cruce RVIE vs COBROS: {len(regs) - len(faltan)} ventas cuadran · "
              f"{len(faltan)} declaradas que NO están en cobros:")
        for r in faltan:
            print(f"   ⚠️ {r['comp']:14} {r['contraparte'][:34]:34} Total S/{r['total']:>8,.2f}")
        if not faltan:
            print("   ✅ toda venta declarada está en nuestro maestro de cobros.")


if __name__ == "__main__":
    main()
