#!/usr/bin/env python3
"""
tools/incorporar_rvie.py — Incorpora las VENTAS del RVIE oficial SUNAT (05. TRIBUTARIO_SIRE)
al maestro COBROS. Destapa las ventas
que efact NO reporta — sobre todo la serie E001 (SUNAT SOL),
que de otro modo queda invisible en cobros.

El RVIE es la fuente OFICIAL de ventas (lo declarado). Reglas:
  · Dedupe por comprobante normalizado (idempotente; no duplica lo ya incorporado a mano).
  · Nota de crédito (Tipo 07 o total < 0): NO entra como fila; ANULA su factura destino, que
    el propio RVIE identifica en el campo "Nro CP Modificado". Ese enlace es del emisor, no una
    heurística → se aplica solo. Si la factura ya estaba cobrada, NO se toca el cobro: se marca
    como devolución para validar. Sin el campo (o destino ausente del maestro) → se reporta para
    revisión manual; jamás se anula adivinando por monto (con gemelas se falla: una NC
    que anulaba una factura se había aplicado a su gemela del mismo importe).
  · Venta a la OTRA empresa del grupo (RUC ALT↔MER) → "Inter-Cía" (se netea en consolidado).
  · Detracción 12% si es servicio > S/700 (nuestras ventas son servicios). Neto = total − det.
  · Cobro: se marca "Cobrado" SOLO con abono que calce (neto o bruto exacto ±1) en la
    conciliación; sin evidencia → Vencido/Por Vencer por antigüedad. Jamás cerrar por adivinar.

Uso:  python3 tools/incorporar_rvie.py --empresa ALT            # todos los meses RVIE disponibles
      python3 tools/incorporar_rvie.py --empresa MER --mes 2026-01,2026-05
"""
from __future__ import annotations

import os
import re
import sys
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import cierre as CI            # noqa: E402
from lib import policy as POL
from lib import sire as SIRE          # noqa: E402
from lib import empresas as EMPX        # noqa: E402
from lib import formato as FMT          # noqa: E402
from lib import sources as S           # noqa: E402

ROOT = CI.ROOT
NCOLS = 15
RUC = {"ALT": "20900000011", "MER": "20900000021"}
HOJA = {"ALT": "ALTAMIRA - Cobros", "MER": "MERAK - Cobros"}
HOJA_CONC = {"ALT": "📊 ALTAMIRA - Movimientos", "MER": "📊 MERAK - Movimientos"}
COL_MONEDA = (5, 7, 8, 11, 12)
COL_FECHA = (2, 10)

FILL_OK = PatternFill("solid", fgColor=FMT.FILL_ESTADO["ok"])
FILL_PEND = PatternFill("solid", fgColor=FMT.FILL_ESTADO["pendiente"])
FILL_NEU = PatternFill("solid", fgColor=FMT.FILL_ESTADO["neutro"])


def _norm(comp) -> str:
    m = re.match(r"([A-Za-z0-9]{4})-0*(\d+)", str(comp or ""))
    return f"{m.group(1)}-{m.group(2)}" if m else str(comp or "")


def _fecha(s):
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(s).strip()[:10], fmt).date()
        except (ValueError, TypeError):
            continue
    return None


# ─── Fuente: xlsx RVIE estructurados ────────────────────────────────────────
def leer_rvie(emp: str, rango) -> list[dict]:
    """Ventas declaradas. El lector vive en `lib/sire.py` (era una de tres copias); aquí solo
    queda el alias `cli`, que es como llama este tool a la contraparte."""
    return [{**d, "cli": d["nombre"]} for d in SIRE.leer(emp, "RVIE", rango)]


# ─── Abonos de la conciliación (evidencia de cobro) ─────────────────────────
def cargar_abonos(emp: str) -> list[dict]:
    ruta = S.latest("conciliacion")
    if not ruta:
        return []
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    if HOJA_CONC[emp] not in wb.sheetnames:
        return []
    out = []
    for row in wb[HOJA_CONC[emp]].iter_rows(min_row=6, values_only=True):
        if not row or len(row) < 5:
            continue
        ab = row[4]
        if not isinstance(ab, (int, float)) or ab <= 0:
            continue
        f = row[0]
        fx = f.date() if isinstance(f, (datetime, date)) else _fecha(f)
        if fx:
            out.append({"f": fx, "m": round(float(ab), 2), "desc": str(row[2] or ""),
                        "ref": str(row[7] or ""), "usado": False})
    return out


def casar_cobro(doc: dict, abonos: list[dict]):
    """Calza el cobro por NETO o BRUTO exacto (±1). Devuelve (abono, modo) o None."""
    if doc["es_nc"] or doc["moneda"] != "PEN" or doc["total"] <= 0 or not doc["femi"]:
        return None
    det = EMPX.detr_esperada(doc["total"])          # ya aplica UMBRAL_SPOT (0 si ≤ S/700)
    objetivos = [(round(doc["total"] - det, 2), "neto"), (round(doc["total"], 2), "bruto")]
    for monto, modo in objetivos:
        for c in abonos:
            if c["usado"] or abs(c["m"] - monto) > 1.0:
                continue
            if not (doc["femi"] - timedelta(days=5) <= c["f"] <= doc["femi"] + timedelta(days=120)):
                continue
            c["usado"] = True
            return c, modo
    return None


# ─── Escritura al maestro COBROS (preservando formato) ──────────────────────
def incorporar(emp: str, docs: list[dict], abonos: list[dict], aplicar: bool | None = None) -> dict:
    # `aplicar` explícito para que los tests puedan ejercitar AMBOS caminos; por
    # defecto lo decide la línea de comandos, que es el contrato del gate.
    if aplicar is None:
        aplicar = "--aplicar" in sys.argv
    base = S.latest("cobros")
    if not base:
        sys.exit("❌ No hay COBROS_*.xlsx base en 04. FINANZAS/")
    out = os.path.join(ROOT, "04. FINANZAS", "08. COBROS", f"COBROS_{date.today():%Y-%m}.xlsx")
    wb = openpyxl.load_workbook(base)
    ws = wb[HOJA[emp]]
    corte = date.today()
    otro = RUC["MER"] if emp == "ALT" else RUC["ALT"]

    ftot = next(r for r in range(1, ws.max_row + 1)
                if str(ws.cell(r, 1).value or "").startswith("TOTALES"))
    existentes, ya = [], set()
    for r in range(4, ftot):
        existentes.append({c: ws.cell(r, c).value for c in range(1, NCOLS + 1)})
        ya.add(_norm(ws.cell(r, 3).value))

    nuevos, saltados, ncs = [], 0, []
    for d in docs:
        # Las notas de crédito NO se agregan como filas: anulan la factura DESTINO, que el
        # propio RVIE identifica en "Nro CP Modificado". Con ese dato se aplican solas; sin él
        # (campo vacío) se REPORTAN para revisión manual — nunca se adivina por monto.
        if d["es_nc"]:
            ncs.append(d)
            continue
        if _norm(d["comp"]) in ya:
            saltados += 1
            continue
        det = EMPX.detr_esperada(d["total"])        # ya aplica UMBRAL_SPOT (0 si ≤ S/700)
        neto = round(d["total"] - det, 2)
        cobro = casar_cobro(d, abonos)
        inter = d["ruc"] == otro
        if inter:
            estado, cobrado, saldo = "Inter-Cía", 0, 0
        elif cobro:
            estado, cobrado, saldo = "Cobrado", cobro[0]["m"], 0
        else:
            venc = d["femi"] and (corte - d["femi"]).days > 30
            estado, cobrado, saldo = ("Vencido" if venc else "Por Vencer"), 0, neto
        obs = ["fuente RVIE oficial SUNAT"]
        if d["comp"].upper().startswith("E001"):
            obs.append("E001 (SUNAT SOL) — no aparece en efact")
        if inter:
            obs.append(f"venta inter-compañía a {'MERAK' if emp == 'ALT' else 'ALTAMIRA'} (netea en consolidado)")
        if cobro:
            obs.append(f"cobro {cobro[1]} {cobro[0]['f']:%d/%m} S/{cobro[0]['m']:,.2f} «{cobro[0]['desc'][:20]}»")
        elif not d["es_nc"] and not inter:
            obs.append("sin abono que calce en el banco — por cobrar")
        nuevos.append({"d": d, "det": det, "neto": neto, "estado": estado,
                       "cobrado": round(cobrado, 2), "saldo": round(saldo, 2),
                       "fcobro": cobro[0]["f"] if cobro else None, "obs": " · ".join(obs)})

    # reconstruir bloque de datos: existentes + nuevos, ordenado por emisión
    filas = []
    for rec in existentes:
        fv = rec[2]
        fk = fv.date() if isinstance(fv, (datetime, date)) else (_fecha(fv) or date.max)
        filas.append(("old", fk, rec))
    for n in nuevos:
        filas.append(("new", n["d"]["femi"] or date.max, n))
    filas.sort(key=lambda x: x[1])

    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= 4:
            ws.unmerge_cells(str(mr))
    for r in range(4, ws.max_row + 2):
        for c in range(1, NCOLS + 1):
            ws.cell(r, c).value = None
            ws.cell(r, c).fill = PatternFill()

    fila = 4
    for tipo, _fk, rec in filas:
        if tipo == "old":
            for c in range(1, NCOLS + 1):
                ws.cell(fila, c).value = rec[c]
            estado = str(rec[13] or "")
        else:
            d, n = rec["d"], rec
            femi = d["femi"]
            ws.cell(fila, 1).value = CI.ABREV_MES[femi.month] if femi else ""
            ws.cell(fila, 2).value = femi
            ws.cell(fila, 3).value = d["comp"]
            cli = d["cli"]
            ws.cell(fila, 4).value = cli.title() if cli.isupper() else cli
            ws.cell(fila, 5).value = round(d["total"], 2)
            ws.cell(fila, 6).value = EMPX.TASA_SPOT if n["det"] else 0
            ws.cell(fila, 7).value = round(n["det"], 2)
            ws.cell(fila, 8).value = n["neto"]
            ws.cell(fila, 9).value = ""
            ws.cell(fila, 10).value = n["fcobro"]
            ws.cell(fila, 11).value = n["cobrado"]
            ws.cell(fila, 12).value = n["saldo"]
            ws.cell(fila, 13).value = n["estado"]
            ws.cell(fila, 14).value = None
            ws.cell(fila, 15).value = n["obs"]
            estado = n["estado"]
        for c in range(1, NCOLS + 1):
            ws.cell(fila, c).font = Font(bold=False)
        for c in COL_MONEDA:
            ws.cell(fila, c).number_format = FMT.FMT_MONEDA
        for c in COL_FECHA:
            ws.cell(fila, c).number_format = FMT.FMT_FECHA
        el = estado.lower()
        fillc = FILL_OK if "cobrado" in el else (FILL_NEU if el in ("anulado", "inter-cía") else FILL_PEND)
        for c in range(1, NCOLS + 1):
            ws.cell(fila, c).fill = fillc
        fila += 1

    # ── Aplicar las NC sobre su factura destino (dato del propio RVIE) ──
    nc_ok, nc_huerf = [], []
    for n in ncs:
        destino = _norm(n["modifica"])
        if not destino:
            nc_huerf.append((n["comp"], "sin 'Nro CP Modificado' en el RVIE"))
            continue
        hallada = False
        for r in range(4, fila):        # `fila` es la primera libre = futura fila de TOTALES
            if _norm(ws.cell(r, 3).value) != destino:
                continue
            hallada = True
            est = str(ws.cell(r, 13).value or "")
            cobrado = CI.num(ws.cell(r, 11).value)
            if est.lower() == "anulado":
                break                                   # ya estaba aplicada (idempotente)
            if cobrado:
                # anular una factura YA COBRADA es una devolución: no se toca el cobro,
                # se marca y se reporta (regla del cobrador).
                obs = str(ws.cell(r, 15).value or "")
                ws.cell(r, 15).value = (f"⚠️ NC {n['comp']} la anula pero estaba COBRADA "
                                        f"S/{cobrado:,.2f} → es DEVOLUCIÓN, validar. " + obs)[:250]
                nc_huerf.append((n["comp"], f"anula {destino} que ya estaba cobrada — devolución"))
            else:
                ws.cell(r, 13).value = "Anulado"
                ws.cell(r, 11).value = 0
                ws.cell(r, 12).value = 0
                ws.cell(r, 15).value = (f"ANULADA por NC {n['comp']} ({n['total']:,.2f}) — "
                                        f"enlace del RVIE, campo 'Nro CP Modificado'")
                nc_ok.append((n["comp"], destino))
            break
        if not hallada:
            nc_huerf.append((n["comp"], f"anula {destino}, que no está en el maestro"))

    # ── TOTALES: activos (excluye Anulado e Inter-Cía) ──
    def _act(r):
        return str(ws.cell(r, 13).value or "").lower() not in ("anulado", "inter-cía")

    tot_row = fila
    ws.cell(tot_row, 1).value = "TOTALES"
    ws.cell(tot_row, 1).font = Font(bold=True)
    for col in COL_MONEDA:
        v = round(sum(CI.num(ws.cell(r, col).value) for r in range(4, tot_row) if _act(r)), 2)
        ws.cell(tot_row, col).value = v
        ws.cell(tot_row, col).number_format = FMT.FMT_MONEDA
        ws.cell(tot_row, col).font = Font(bold=True)
    ws.cell(tot_row, 15).value = "Activos (Anulado e Inter-Cía excluidos)"
    for c in range(1, NCOLS + 1):
        ws.cell(tot_row, c).fill = PatternFill("solid", fgColor=FMT.TOTAL_FILL)
    FMT.pulir_hoja(ws)

    # El gate mide el DELTA, no la hoja entera. Medir el total restatado hacía que el umbral
    # disparara en cada corrida pasara lo que pasara — y un control que siempre pide permiso
    # enseña a concederlo sin mirar, que es peor que no tenerlo.
    d_nuevos = [n["d"]["total"] for n in nuevos]
    if not POL.exigir(POL.evaluar(POL.Accion(
            tool="incorporar_rvie", maestro=out, empresa=None, capacidad="cobros.escribir",
            periodo=POL.periodo_de(out), aplicar=aplicar,
            monto=max(d_nuevos, default=0.0), total=round(sum(d_nuevos), 2),
            evidencia="RVIE oficial de SUNAT (ventas)",
            forzar="--forzar" in sys.argv))):
        return {"preview": True, "nuevos_n": len(nuevos),
                "nuevos_total": round(sum(d_nuevos), 2),
                "detalle": [(n["d"]["comp"], n["d"]["femi"], (n["d"].get("nombre") or n["d"].get("cli") or n["d"].get("nom") or "")[:34],
                             n["d"]["total"]) for n in nuevos]}
    wb.save(out)
    return {
        "out": out, "filas": tot_row - 4, "nuevos": len(nuevos), "saltados": saltados,
        "cobrado_n": sum(1 for n in nuevos if n["estado"] == "Cobrado"),
        "e001": sum(1 for n in nuevos if n["d"]["comp"].upper().startswith("E001")),
        "nc_ok": nc_ok, "nc_huerf": nc_huerf,
        "inter": sum(1 for n in nuevos if n["estado"] == "Inter-Cía"),
        "facturado": round(sum(n["d"]["total"] for n in nuevos
                               if n["estado"] != "Inter-Cía"), 2),
    }


def _rango():
    if "--mes" not in sys.argv:
        return None
    val = sys.argv[sys.argv.index("--mes") + 1]
    a, b = (val.split(",") + [val])[:2]
    return (a.strip(), b.strip())


def main():
    emp = (sys.argv[sys.argv.index("--empresa") + 1].upper()
           if "--empresa" in sys.argv else "ALT")
    docs = leer_rvie(emp, _rango())
    if not docs:
        sys.exit("❌ No hay ventas RVIE para ese rango.")
    r = incorporar(emp, docs, cargar_abonos(emp))
    if r.get("preview"):
        print(f"🔍 Preview ({r['nuevos_n']} comprobantes nuevos, "
              f"S/{r['nuevos_total']:,.2f}) — nada escrito:")
        for comp, fe, nom, tot in r["detalle"]:
            print(f"     {comp:16} {fe} {nom:34} {tot:>11,.2f}")
        print("\n   Para escribirlo:  .venv/bin/python3 tools/incorporar_rvie.py "
              "--empresa %s --aplicar" % emp)
        return
    print("✅ Generado:", os.path.basename(r["out"]))
    print(f"  {emp}: {r['nuevos']} ventas nuevas ({r['saltados']} ya estaban) · "
          f"{r['e001']} E001 · {r['cobrado_n']} casadas al banco · {r['inter']} inter-cía")
    print(f"  Facturación nueva incorporada (activa): S/{r['facturado']:,.2f}")
    if r["nc_ok"]:
        print(f"  ✔ {len(r['nc_ok'])} NC aplicadas automáticamente (enlace del RVIE):")
        for nc, dst in r["nc_ok"]:
            print(f"     {nc} → anuló {dst}")
    if r["nc_huerf"]:
        print(f"  ⚠️ {len(r['nc_huerf'])} NC que requieren revisión manual:")
        for nc, motivo in r["nc_huerf"]:
            print(f"     {nc}: {motivo}")


if __name__ == "__main__":
    main()
