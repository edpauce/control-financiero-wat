"""
tools/lib/sire.py — Lector único de los RVIE/RCE estructurados (`05. TRIBUTARIO_SIRE`).

El RVIE (ventas) y el RCE (compras) del SIRE son la **fuente oficial**: lo que SUNAT tiene
declarado. Los maestros COBROS y PROVEEDORES son nuestro trabajo. Que las dos cosas se puedan
comparar es el punto entero del cruce tributario.

Por qué este módulo existe: había TRES copias del mismo lector (`incorporar_rvie`,
`incorporar_rce`, `detraccion_compras`), con nombres de campo distintos para la misma columna
y filtros de calidad en unas sí y en otras no. Una regla que solo vive en una copia no es una
regla. Aquí está una vez.

Layout de los xlsx (los escribe `ingerir_sire_txt.py`; fila 1 título, fila 2 cabecera):
  1 Comprobante · 2 Tipo · 3 Fecha · 4 RUC contraparte · 5 Contraparte
  6 BI Gravada · 7 IGV · 8 Total · 9 Moneda · 10 Anula/Modifica
"""
from __future__ import annotations

import datetime as dt
import glob
import os
import re

import openpyxl

from . import cierre as CI

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DIR = {"ALT": "01. ALTAMIRA/05. TRIBUTARIO_SIRE/2026",
       "MER": "02. MERAK/05. TRIBUTARIO_SIRE/2026"}
ABREV = {"ALTAMIRA": "ALT", "MERAK": "MER"}
RUC_BCP = "20100047218"


def _abrev(emp: str) -> str:
    return ABREV.get(emp.upper(), emp.upper()[:3])


def _fecha(s):
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(str(s).strip()[:10], fmt).date()
        except (ValueError, TypeError):
            continue
    return None


def archivos(emp: str, libro: str) -> dict[str, str]:
    """{'2026-05': ruta} de los xlsx disponibles. `libro` es 'RVIE' o 'RCE'.

    Solo los que traen periodo en el nombre: puede existir un `SIRE_RVIE_ALT_sin-periodo.xlsx` cuyo
    contenido no se puede imputar a un mes, y meterlo en un cruce mensual inventa un descuadre.
    """
    ab = _abrev(emp)
    out = {}
    for f in sorted(glob.glob(os.path.join(RAIZ, DIR[ab], f"SIRE_{libro}_{ab}_2026-*.xlsx"))):
        m = re.search(r"(2026-\d\d)", os.path.basename(f))
        if m:
            out[m.group(1)] = f
    return out


def leer(emp: str, libro: str, rango: tuple[str, str] | None = None) -> list[dict]:
    """Comprobantes del libro oficial. Campos canónicos; `nombre` es la contraparte
    (cliente en el RVIE, proveedor en el RCE)."""
    docs = []
    for periodo, f in sorted(archivos(emp, libro).items()):
        if rango and not (rango[0] <= periodo <= rango[1]):
            continue
        ws = openpyxl.load_workbook(f, data_only=True).active
        for r in range(3, ws.max_row + 1):
            comp = ws.cell(r, 1).value
            if not comp or str(comp).startswith("TOTAL"):
                continue
            ruc = str(ws.cell(r, 4).value or "")
            if not ruc.isdigit():
                # Artefacto de parseo del SIRE (a veces trae una fila
                # '0.00-0.00 / 1.000 / PEN' con RUC='PEN'). Los RUC del SIRE son numéricos
                # siempre; una fila sin RUC no es un comprobante.
                continue
            tipo = str(ws.cell(r, 2).value or "Factura")
            total = CI.num(ws.cell(r, 8).value)
            docs.append({
                "periodo": periodo, "comp": str(comp), "tipo": tipo,
                "femi": _fecha(ws.cell(r, 3).value), "ruc": ruc,
                "nombre": str(ws.cell(r, 5).value or ""),
                "bi": CI.num(ws.cell(r, 6).value), "igv": CI.num(ws.cell(r, 7).value),
                "total": total, "moneda": ws.cell(r, 9).value or "PEN",
                "es_nc": total < 0 or tipo in ("07", "7", "NC"),
                # col 10: el comprobante que esta NC modifica, según el propio SIRE. Con esto
                # la NC se aplica sola y no hay que adivinar por monto (falla con gemelas).
                "modifica": str(ws.cell(r, 10).value or "").strip(),
                "banco": ruc == RUC_BCP,
            })
    return docs


def normalizar(comp: str) -> str:
    """«FA02-08000945» y «FA02-8000945» son el MISMO comprobante.

    El SIRE y los maestros discrepan en los ceros de relleno (y `E001-00000009` convive con
    `E001-9`). Sin normalizar, un mismo documento aparece a la vez como «falta en el maestro»
    y como «no declarado»: dos hallazgos falsos que tapan los verdaderos.
    """
    partes = str(comp).strip().upper().split("-", 1)
    if len(partes) != 2:
        return str(comp).strip().upper()
    serie, num = partes
    return f"{serie}-{num.lstrip('0') or '0'}" if num.isdigit() else f"{serie}-{num}"


def por_periodo(emp: str, libro: str) -> dict[str, float]:
    """{'2026-05': total DECLARADO en ese periodo}. Las NC entran con su signo."""
    tot: dict[str, float] = {}
    for d in leer(emp, libro):
        tot[d["periodo"]] = round(tot.get(d["periodo"], 0.0) + d["total"], 2)
    return tot


def por_emision(emp: str, libro: str, periodos=None) -> dict[str, float]:
    """{'2026-05': total de lo EMITIDO ese mes}, aunque se haya declarado en otro.

    Existe porque el SIRE agrupa por periodo declarado y los maestros por fecha de emisión, y
    una factura puede declararse en un periodo posterior al de su emisión. Comparar un bucket contra el otro acusa un descuadre donde
    solo hubo demora. Para cruzar con un maestro hay que agrupar como agrupa el maestro.

    `periodos` acota qué archivos se miran (la racha contigua): un comprobante declarado en un
    periodo que no tenemos no debe entrar en el total esperado.
    """
    tot: dict[str, float] = {}
    for d in leer(emp, libro):
        if periodos is not None and d["periodo"] not in periodos:
            continue
        if d["femi"] is None:
            continue
        k = f"{d['femi']:%Y-%m}"
        tot[k] = round(tot.get(k, 0.0) + d["total"], 2)
    return tot
