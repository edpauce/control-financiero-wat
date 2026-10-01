"""
tools/lib/cierre.py — Helpers comunes del cierre mensual + carga de config por mes.

El cierre está parametrizado por PERIODO ("AAAA-MM"). Los datos específicos de cada mes
(saldos BN, overrides de categoría/cobros confirmados, notas) viven en tools/config/{PERIODO}.py;
este módulo los carga con defaults seguros y deriva rutas de archivos según la convención
documentada en references/fuentes-datos.md.

Uso típico en un tool:
    from lib import cierre as CI
    cfg = CI.cargar_config("2026-07")
    R = CI.rutas(cfg)                    # rutas de fuentes crudas del mes
    src = CI.first(R["bcp"]["MERAK"])    # archivo más reciente del glob
"""
from __future__ import annotations
import csv
import glob
import importlib.util
import os
import re
import shutil
import sys
from copy import copy as _copy
from pathlib import Path

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
CONFIG_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "config"))

MESES_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
ABREV_MES = {1: "Ene", 2: "Feb", 3: "Mar", 4: "Abr", 5: "May", 6: "Jun",
             7: "Jul", 8: "Ago", 9: "Set", 10: "Oct", 11: "Nov", 12: "Dic"}
CARPETA_MES = {1: "01.ENERO", 2: "02.FEBRERO", 3: "03.MARZO", 4: "04.ABRIL", 5: "05.MAYO",
               6: "06.JUNIO", 7: "07.JULIO", 8: "08.AGOSTO", 9: "09.SEPTIEMBRE",
               10: "10.OCTUBRE", 11: "11.NOVIEMBRE", 12: "12.DICIEMBRE"}
DIR_EMP = {"ALTAMIRA": "01. ALTAMIRA", "MERAK": "02. MERAK"}
ABREV_EMP = {"ALTAMIRA": "ALT", "MERAK": "MER"}

# Ajustes históricos PERMANENTES al flujo/P&L (hechos puntuales, no dependen del mes que se cierra):
# (empresa, "AAAA-MM") → monto a excluir de abonos Y cargos (neutro en saldo).
# La llave lleva el AÑO: con solo el nombre del mes, el ajuste se re-aplicaría cada junio.
AJUSTE_FLUJO = {("MERAK", "2026-06"): 50000.0}   # valor de ejemplo: ida y vuelta de fondos propios, neutro


def periodo_de(anio: int, mes_nombre: str) -> str:
    """('2026', 'Junio') → '2026-06'."""
    return f"{anio}-{MESES_ES.index(mes_nombre) + 1:02d}"

# Campos opcionales de la config por mes y sus defaults
_DEFAULTS = {
    "CERRADO": False,         # True = mes cerrado: los maestros llevan formato manual, NO regenerar
    "CORTE": None,            # "DD/MM" al que llega la data (None → último día conocido)
    "REV_FECHA": None,        # fecha de la revisión de salida (None → hoy)
    "FECHA_AGING": None,      # fecha de corte del aging de CxC (None → deriva de CORTE)
    "SALDO_BN": None,         # {"MERAK": x, "ALTAMIRA": y} al corte — OBLIGATORIO p/ detracciones
    "OP_OVERRIDES": {},       # op BCP → categoría (glosas crípticas identificadas por el usuario)
    "OP_OBS": {},             # op BCP → observación
    "OP_CONTRAPARTE": {},     # op BCP → nombre de la contraparte (REFERENCIA) confirmado a mano
    "INTERCIA_OPS_ALT": set(),# ops de 'DE OTRA CUENTA' en ALT que son inter-cía MERAK→ALT
    "GLOSA_OVERRIDES": [],    # [(keywords_tuple, categoría, observación)] aplicados por glosa;
                              # se fusionan con GLOSA_PERMANENTES de tools/config/glosario.py
    "CONFIRMADOS": {},        # comprobante → (fecha, monto) cobros confirmados por el usuario
    "COBRADO_MANUAL": {},     # fragmento de cliente → fecha (facturas previas confirmadas pagadas)
    "ANULAR_MANUAL": [],      # comprobantes de facturas previas a anular
    "AUTODEP_DOCS": set(),    # docs BN que son auto-depósito (rosado en hojas 04/05)
    "NOTA_09": "",            # texto del encabezado de '09 Pendientes' (vacío → no tocar)
}


def _corte_argv(mes: int) -> str | None:
    """
    Override --corte DD/MM desde la línea de comandos: habilita cierres QUINCENALES de primera
    clase (p.ej. --mes 2026-07 --corte 15/07). Valida que el mes del corte sea el del periodo.
    Con el corte declarado, la compuerta de ingesta exige cobertura solo hasta esa fecha: una
    quincena no necesita --permitir-incompleto.
    """
    if "--corte" not in sys.argv:
        return None
    i = sys.argv.index("--corte")
    if i + 1 >= len(sys.argv):
        sys.exit("❌ --corte necesita un valor DD/MM (p.ej. --corte 15/07)")
    v = sys.argv[i + 1]
    m = re.match(r"^(\d{1,2})/(\d{1,2})$", v)
    if not m:
        sys.exit(f"❌ --corte inválido: {v!r} (formato DD/MM, p.ej. 15/07)")
    dd, mm = int(m.group(1)), int(m.group(2))
    if mm != mes:
        sys.exit(f"❌ --corte {v}: el mes del corte ({mm:02d}) no es el del periodo ({mes:02d})")
    if not 1 <= dd <= 31:
        sys.exit(f"❌ --corte {v}: día fuera de rango")
    return f"{dd:02d}/{mm:02d}"


def glosa_permanentes() -> list:
    """GLOSA_PERMANENTES de tools/config/glosario.py (la lista vacía si no existe el archivo)."""
    ruta = os.path.join(CONFIG_DIR, "glosario.py")
    if not os.path.exists(ruta):
        return []
    spec = importlib.util.spec_from_file_location("glosario_cfg", ruta)
    g = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(g)
    return [tuple(r) for r in getattr(g, "GLOSA_PERMANENTES", [])]


def fusionar_glosas(permanentes: list, del_mes: list) -> list:
    """
    Reglas de glosa efectivas del mes: PRIMERO las del mes (mandan y se evalúan antes), luego las
    permanentes que no comparten ninguna palabra clave con una regla del mes. Las palabras clave
    se comparan en mayúsculas, que es como se aplican sobre la glosa del banco.
    """
    del_mes = [tuple(r) for r in (del_mes or [])]
    usadas = {str(k).upper() for r in del_mes for k in r[0]}
    return del_mes + [r for r in permanentes if not any(str(k).upper() in usadas for k in r[0])]


def cargar_config(periodo: str):
    """Carga tools/config/{periodo}.py, valida y completa defaults + campos derivados."""
    if not re.match(r"^\d{4}-\d{2}$", str(periodo)):
        sys.exit(f"❌ Periodo inválido: {periodo!r} (formato AAAA-MM, p.ej. 2026-07)")
    path = os.path.join(CONFIG_DIR, f"{periodo}.py")
    if not os.path.exists(path):
        sys.exit(f"❌ No existe tools/config/{periodo}.py — cópialo de tools/config/2026-07.py "
                 f"(plantilla) y completa los datos del mes.")
    spec = importlib.util.spec_from_file_location(f"cfg_{periodo.replace('-', '_')}", path)
    cfg = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(cfg)
    for k, v in _DEFAULTS.items():
        if not hasattr(cfg, k):
            setattr(cfg, k, v)
    cfg.GLOSA_OVERRIDES = fusionar_glosas(glosa_permanentes(), cfg.GLOSA_OVERRIDES)
    # derivados
    cfg.PERIODO = periodo
    cfg.ANIO, cfg.MES = int(periodo[:4]), int(periodo[5:7])
    cfg.MES_NOMBRE = MESES_ES[cfg.MES - 1]
    cfg.MES_ABREV = ABREV_MES[cfg.MES]
    pa, pm = (cfg.ANIO, cfg.MES - 1) if cfg.MES > 1 else (cfg.ANIO - 1, 12)
    cfg.PREV_PERIODO = f"{pa}-{pm:02d}"
    cfg.PREV_NOMBRE = MESES_ES[pm - 1]
    corte_cli = _corte_argv(cfg.MES)
    if corte_cli:
        cfg.CORTE = corte_cli                     # cierre quincenal/provisional declarado
        cfg.FECHA_AGING = None                    # re-derivar del corte declarado
    elif not cfg.CORTE:
        cfg.CORTE = f"{ultimo_dia(cfg.ANIO, cfg.MES):02d}/{cfg.MES:02d}"
    # Aging de cuentas por cobrar: se mide al CORTE del mes (el COBROS_{mes} es una foto del cierre,
    # no un documento que envejece solo). Se puede forzar con FECHA_AGING = "AAAA-MM-DD".
    from datetime import date as _date
    if cfg.FECHA_AGING is None:
        dd, mm = (int(x) for x in cfg.CORTE.split("/"))
        cfg.FECHA_AGING = _date(cfg.ANIO, mm, dd)
    elif isinstance(cfg.FECHA_AGING, str):
        cfg.FECHA_AGING = _date.fromisoformat(cfg.FECHA_AGING)
    return cfg


def exigir_abierto(cfg, tool: str = "") -> None:
    """
    Aborta si el mes está CERRADO (tools/config/{mes}.py: CERRADO = True) y no se pasó --forzar.

    Un mes cerrado lleva formato manual en los maestros (colores puestos a mano) y sus
    correcciones se hacen EDITANDO CELDAS en sitio. Regenerarlo pisa ese trabajo. Antes esto era
    solo un comentario "⛔ NO REGENERAR" en la config; ahora el pipeline lo hace cumplir.
    """
    if not getattr(cfg, "CERRADO", False) or "--forzar" in sys.argv:
        return
    quien = f" ({tool})" if tool else ""
    sys.exit(f"⛔ {cfg.MES_NOMBRE} {cfg.ANIO} está CERRADO (tools/config/{cfg.PERIODO}.py: CERRADO = True)"
             f"{quien}. Regenerarlo pisa el formato manual de los maestros: las correcciones van\n"
             f"   en sitio, editando celdas. Si de verdad quieres regenerar el mes, añade --forzar.")


def ultimo_dia(anio: int, mes: int) -> int:
    import calendar
    return calendar.monthrange(anio, mes)[1]


def config_minima(periodo: str):
    """
    Config sintética para tools que solo necesitan el CALENDARIO del periodo (validar fuentes,
    enriquecer meses viejos), sin exigir tools/config/{periodo}.py — ese archivo trae los datos
    del cierre (SALDO_BN, overrides) y se crea DESPUÉS de descargar y validar las fuentes.
    """
    if not re.match(r"^\d{4}-\d{2}$", str(periodo)):
        sys.exit(f"❌ Periodo inválido: {periodo!r} (formato AAAA-MM, p.ej. 2026-07)")
    cfg = type("CfgMin", (), {})()
    for k, v in _DEFAULTS.items():
        setattr(cfg, k, v)
    cfg.PERIODO = periodo
    cfg.ANIO, cfg.MES = int(periodo[:4]), int(periodo[5:7])
    if not 1 <= cfg.MES <= 12:
        sys.exit(f"❌ Periodo inválido: {periodo!r} (mes fuera de 01..12)")
    cfg.MES_NOMBRE = MESES_ES[cfg.MES - 1]
    cfg.MES_ABREV = ABREV_MES[cfg.MES]
    pa, pm = (cfg.ANIO, cfg.MES - 1) if cfg.MES > 1 else (cfg.ANIO - 1, 12)
    cfg.PREV_PERIODO = f"{pa}-{pm:02d}"
    cfg.PREV_NOMBRE = MESES_ES[pm - 1]
    cfg.CORTE = _corte_argv(cfg.MES) or f"{ultimo_dia(cfg.ANIO, cfg.MES):02d}/{cfg.MES:02d}"
    from datetime import date as _date
    dd = int(cfg.CORTE[:2])
    cfg.FECHA_AGING = _date(cfg.ANIO, cfg.MES, dd)
    return cfg


def rutas(cfg) -> dict:
    """Globs de las fuentes crudas del mes por empresa, según la convención de archivado."""
    p, a = cfg.PERIODO, cfg.ANIO
    r = {"bcp": {}, "bcp_ult": {}, "bn": {}, "sunat": {}, "efact": {}}
    for emp in ("ALTAMIRA", "MERAK"):
        d, ab = DIR_EMP[emp], ABREV_EMP[emp]
        # .csv, .txt o .xlsx: el export xlsx trae columnas extra, el parser resuelve por encabezado.
        r["bcp"][emp] = f"{d}/01. EECC/{a}/EECC_{ab}_{p}_movimientos.*"
        r["bcp_ult"][emp] = f"{d}/01. EECC/{a}/EECC_{ab}_{p}_UltimosMov.*"
        r["bn"][emp] = f"{d}/03. DETRACCIONES/{a}/01. EECC_BN/BN_{ab}_{p}_movimientos.*"
        r["sunat"][emp] = f"{d}/03. DETRACCIONES/{a}/02. Constancias_SUNAT/DETR-SUNAT_{ab}_{p}.csv"
        # Comprobantes emitidos (efact): el archivo operativo vive en 02. FACTURACION (convención
        # histórica del proyecto), NO en 05. TRIBUTARIO.
        r["efact"][emp] = f"{d}/02. FACTURACION/{a}/Emitidas/FACT-EMI_{ab}_{p}_Reporte.xlsx"
    return r


def first(pat: str) -> str | None:
    """Archivo más reciente que coincide con el glob (relativo a la raíz). Ignora temporales."""
    m = [f for f in glob.glob(os.path.join(ROOT, pat), recursive=True)
         if not os.path.basename(f).startswith(("~$", ".~lock"))]
    return max(m, key=os.path.getmtime) if m else None


def respaldar(ruta, etiqueta: str | None = None) -> Path:
    """
    Copia `ruta` a `{ruta}.{etiqueta o AAAA-MM-DD}.bak` ANTES de escribir un maestro en sitio.

    Un slot por etiqueta (o por día): si el .bak de hoy ya existe, NO se pisa. El primer
    respaldo del día es el estado previo a cualquier corrida — justo el que hace falta cuando
    la segunda corrida rompe algo (un único slot `.bak` se destruía a sí mismo en la segunda
    corrida). Devuelve la ruta del .bak, exista ya o recién creado. Se llama antes del
    `load_workbook` de cada tool que guarda sobre el mismo archivo.
    """
    from datetime import date as _date
    src = Path(ruta)
    bak = Path(f"{src}.{etiqueta or _date.today().isoformat()}.bak")
    if not bak.exists():
        shutil.copy2(src, bak)
    return bak


def num(x) -> float:
    try:
        return float(str(x).replace(",", "")) if str(x).strip() not in ("", "None") else 0.0
    except (ValueError, TypeError):
        return 0.0


def fkey(f) -> tuple:
    """Clave ordenable (aaaa,mm,dd) desde 'DD/MM/YYYY'. Normaliza a 2 dígitos: sin esto,
    ('2026','6','1') > ('2026','06','30') y un movimiento del 1/jun se trataría como más nuevo
    que el 30/jun (doble conteo al fusionar histórico + últimos movimientos)."""
    d, m, y = str(f)[:10].split("/")
    return (f"{int(y):04d}", f"{int(m):02d}", f"{int(d):02d}")


def comp_norm(s) -> str | None:
    """Normaliza un comprobante a 'F001-123' (acepta E001/F001, con o sin ceros)."""
    m = re.search(r'([EF]\d{3})-?0*(\d+)', str(s))
    return f"{m.group(1)}-{int(m.group(2))}" if m else None


def _fecha_ddmmaaaa(v) -> str | None:
    """Normaliza la celda de fecha a 'DD/MM/AAAA'. Acepta texto o datetime/date (Excel guarda
    las fechas como datetime, no como texto)."""
    from datetime import date as _d, datetime as _dt
    if isinstance(v, (_dt, _d)):
        return v.strftime("%d/%m/%Y")
    s = str(v or "").strip()
    return s if re.match(r"^\d{1,2}/\d{1,2}/\d{4}", s) else None


def leer_ultimos(path: str) -> list[dict]:
    """Lee 'últimos movimientos' BCP (.csv o .xlsx) → [{fecha, desc, monto, operacion}]."""
    out = []
    if path.lower().endswith(".csv"):
        with open(path, "rb") as fh:
            texto = fh.read().decode("latin-1", "replace")
        rows = list(csv.reader(texto.splitlines()))
        hj = next((i for i, r in enumerate(rows) if r and str(r[0]).strip() == "Fecha"), None)
        src = rows[hj + 1:] if hj is not None else []
    else:
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True).active
        src = [list(r) for r in ws.iter_rows(values_only=True)]
        hj = next((i for i, r in enumerate(src) if r and str(r[0]).strip() == "Fecha"), None)
        src = src[hj + 1:] if hj is not None else []
    for r in src:
        if not r:
            continue
        f = _fecha_ddmmaaaa(r[0])
        if not f:
            continue
        r = (list(r) + [""] * 6)[:6]
        out.append({"fecha": f, "desc": str(r[2] or "").strip(),
                    "monto": num(r[3]), "operacion": str(r[5] or "").strip()})
    # Un archivo con filas de datos pero 0 movimientos parseados significa que el layout cambió.
    # Devolver [] en silencio borra los últimos días del mes del cierre: hay que gritar.
    if src and not out:
        sys.exit(f"❌ {os.path.basename(path)}: {len(src)} filas tras el encabezado pero 0 "
                 f"movimientos parseados — cambió el formato del export. Revisa el archivo.")
    return out


def verificar_cobertura(movs: list, cfg, empresa: str, fuente: str = "") -> None:
    """
    Aborta si los movimientos no llegan al CORTE del mes.

    El BCP publica el último día del mes recién al día siguiente: si el histórico se descarga
    el 30/06, llega solo al 29/06. Una vez eso hizo desaparecer los movimientos del último día
    y hubo que reconstruirlos a mano desde el EECC oficial.
    Escape: --permitir-incompleto (para cierres provisionales a mitad de mes).
    """
    if not movs:
        sys.exit(f"❌ {empresa}: 0 movimientos en {fuente or 'el EECC'} para {cfg.PERIODO}.")
    ultimo = max(fkey(m["fecha"]) for m in movs)
    corte = fkey(f"{cfg.CORTE}/{cfg.ANIO}")
    if ultimo >= corte or "--permitir-incompleto" in sys.argv:
        return
    d, m, a = ultimo[2], ultimo[1], ultimo[0]
    sys.exit(
        f"❌ {empresa} {cfg.PERIODO}: los movimientos llegan al {d}/{m}/{a}, no al {cfg.CORTE}.\n"
        f"   El BCP publica el último día recién al día siguiente. Vuelve a descargar el EECC\n"
        f"   con el mes ya cerrado (workflows/descargar_portales.md).\n"
        f"   Si el cierre es provisional y lo sabes, añade --permitir-incompleto.")


def copiar_estilo(ws, src_row: int, dst_row: int, ncols: int = 9):
    for c in range(1, ncols + 1):
        s, d = ws.cell(src_row, c), ws.cell(dst_row, c)
        d.font = _copy(s.font); d.border = _copy(s.border); d.fill = _copy(s.fill)
        d.number_format = s.number_format; d.alignment = _copy(s.alignment)


def arg_mes(default: str | None = None) -> str:
    """Lee --mes AAAA-MM de sys.argv (o default). Sale con error si falta."""
    if "--mes" in sys.argv:
        i = sys.argv.index("--mes")
        if i + 1 >= len(sys.argv):
            sys.exit("❌ --mes necesita un valor AAAA-MM (p.ej. --mes 2026-07)")
        return sys.argv[i + 1]
    if default:
        return default
    sys.exit("❌ Falta --mes AAAA-MM (p.ej. --mes 2026-07)")
