"""
tools/lib/formato.py — El estándar visual de la casa, en UN solo lugar.

Antes cada writer improvisaba su paleta (el mismo azul 1F4E79 repetido en 4 tools, hojas sin
bordes, freezes rotos en mitad de tabla). Todo lo que un maestro necesita para verse impecable
y validarse de un vistazo sale de aquí:

  · PALETA / FILL_ESTADO ... colores semánticos (verde=ok, rojo=pendiente, ámbar=revisar)
  · pulir_hoja(ws) ......... normaliza CUALQUIER hoja: formato de dinero y fecha por encabezado,
                             anchos por contenido, freeze bajo el encabezado real, bordes finos.
                             NUNCA toca valores ni fills existentes (los colores por estado son
                             información, no decoración).

Regla de oro heredada del incidente de los merges: el formato lo produce el TOOL, no la mano.
Un retoque manual sobre una hoja regenerable muere en la siguiente corrida.
"""
from __future__ import annotations

import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─── Paleta de la casa ───────────────────────────────────────────────────────
NAVY = "1F4E79"          # encabezados de tabla
NAVY_SUAVE = "D9E1F2"    # sub-encabezados / totales intermedios
GRIS_BANDA = "44546A"    # banners de mes (movimientos)
TOTAL_FILL = "B4C7E7"    # filas TOTAL
NOTA_GRIS = "808080"     # notas al pie

FILL_ESTADO = {          # semántica compartida por todas las hojas
    "ok": "E2EFDA",      # ✅ verde
    "pendiente": "F8CBAD",   # 🔴 rojo suave
    "verificar": "FFF2CC",   # 🟡 ámbar
    "propuesta": "FDE9D9",   # 🟠 naranja suave (P3b: por confirmar)
    "info": "DDEBF7",        # 🔵 azul suave
    "neutro": "F2F2F2",      # 🔘 gris
    "anulado": "D9D9D9",
}

FMT_MONEDA = '#,##0.00'
FMT_FECHA = 'dd/mm/yyyy'
BORDE_FINO = Border(*[Side("thin", color="D9D9D9")] * 4)

# Encabezados que identifican una columna de DINERO o de FECHA (match por substring, sin tildes)
_COL_MONEDA = ("S/", "MONTO", "TOTAL", "ABONO", "CARGO", "SALDO", "IMPORTE", "COBRADO",
               "PENDIENTE", "CUOTA", "DETR", "ESPERADA", "DEPOSITADA", "DECLARADA", "NETO",
               "CAPITAL", "VENTAS", "INGRESO", "EGRESO", "PAGADO", "VARIACION", "VARIACIÓN")
_COL_FECHA = ("FECHA", "EMISION", "EMISIÓN", "VENCIM", "F. DEP", "F.COBRO", "F. COBRO",
              "PROX", "PRÓX")
_COL_NO_MONEDA = ("Nº", "N°", "NRO", "CONSTANCIA", "RUC", "OPERACI", "DOCUMENTO", "%")


def encabezado(ws, fila: int, ncols: int, relleno: str = NAVY) -> None:
    """Estilo de encabezado de tabla: negrita blanca sobre azul, centrado, con wrap."""
    for c in range(1, ncols + 1):
        cell = ws.cell(fila, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=relleno)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fila_encabezado(ws) -> int | None:
    """La fila de encabezado real: la primera (entre las 10 primeras) con ≥3 textos cortos."""
    for r in range(1, min(ws.max_row, 10) + 1):
        textos = [c for c in ws[r] if isinstance(c.value, str) and 0 < len(c.value.strip()) <= 40]
        if len(textos) >= 3:
            return r
    return None


def _es(colname: str, claves: tuple) -> bool:
    up = str(colname or "").upper()
    return any(k in up for k in claves)


def pulir_hoja(ws, con_bordes: bool = True) -> dict:
    """
    Normaliza el formato de una hoja SIN tocar valores, merges ni fills existentes.
    Devuelve un resumen de lo que cambió (para el reporte del pulidor).
    """
    import datetime as _dt
    cambios = {"moneda": 0, "fecha": 0, "anchos": 0, "freeze": None, "bordes": 0}
    hdr = _fila_encabezado(ws)
    if hdr is None:
        return cambios

    # 1) freeze: bajo el encabezado. Corrige los rotos (congelados a mitad de tabla) y los
    #    ausentes; respeta uno existente que ya esté a la altura del encabezado o antes.
    objetivo = f"A{hdr + 1}"
    actual = ws.freeze_panes
    fila_actual = int(re.sub(r"\D", "", actual)) if actual and re.sub(r"\D", "", actual) else None
    if actual is None or (fila_actual is not None and fila_actual > hdr + 1):
        ws.freeze_panes = objetivo
        cambios["freeze"] = objetivo

    # 2) tipo de cada columna, por su encabezado
    tipos = {}
    for c in range(1, ws.max_column + 1):
        nombre = str(ws.cell(hdr, c).value or "")
        if _es(nombre, _COL_NO_MONEDA):
            tipos[c] = None
        elif _es(nombre, _COL_MONEDA):
            tipos[c] = "moneda"
        elif _es(nombre, _COL_FECHA):
            tipos[c] = "fecha"

    # 3) formatos en la zona de datos
    for r in range(hdr + 1, ws.max_row + 1):
        for c, tipo in tipos.items():
            cell = ws.cell(r, c)
            if cell.value is None or tipo is None:
                continue
            if tipo == "moneda" and isinstance(cell.value, (int, float)) \
                    and cell.number_format in ("General", "@"):
                cell.number_format = FMT_MONEDA
                cambios["moneda"] += 1
            elif tipo == "fecha" and isinstance(cell.value, (_dt.datetime, _dt.date)) \
                    and cell.number_format in ("General", "@", "yyyy-mm-dd",
                                               "yyyy-mm-dd h:mm:ss", "yyyy\\-mm\\-dd",
                                               "yyyy\\-mm\\-dd\\ h:mm:ss"):
                # solo los defaults (General y el ISO que openpyxl pone solo); un formato
                # de fecha puesto a mano se respeta
                cell.number_format = FMT_FECHA
                cambios["fecha"] += 1

    # 4) anchos: solo columnas SIN ancho deliberado. Ojo con openpyxl: INDEXAR
    #    column_dimensions crea la entrada con width=13.0 y customWidth=True, así que ni width
    #    ni customWidth discriminan — lo que discrimina es si la letra YA EXISTÍA.
    for c in range(1, ws.max_column + 1):
        letra = get_column_letter(c)
        if letra in ws.column_dimensions and ws.column_dimensions[letra].width:
            continue
        largo = 0
        for r in range(1, min(ws.max_row, 200) + 1):
            v = ws.cell(r, c).value
            if v is not None and r != 1:            # los títulos largos de A1 no dictan el ancho
                largo = max(largo, len(str(v)))
        if largo:
            ws.column_dimensions[letra].width = min(max(largo + 2, 9), 44)
            cambios["anchos"] += 1

    # 5) bordes finos en la zona de datos, solo donde no hay ninguno
    if con_bordes:
        for r in range(hdr, ws.max_row + 1):
            fila = ws[r]
            if not any(c.value is not None for c in fila):
                continue
            for c in fila[:ws.max_column]:
                b = c.border
                if not (b and ((b.left and b.left.style) or (b.top and b.top.style))):
                    c.border = BORDE_FINO
                    cambios["bordes"] += 1
    return cambios
