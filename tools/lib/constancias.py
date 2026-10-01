#!/usr/bin/env python3
"""
tools/lib/constancias.py — Las constancias SUNAT como PRUEBA DE COBRO.

Por qué existe: la cartera declarada incluía facturas que ya no eran deuda. Una
detracción depositada demuestra que la factura se cobró — el 12% solo se deposita cuando hay pago,
lo haga el cliente (paga neto y deposita) o nosotros (cobramos bruto y auto-detraemos). El maestro
de cobros nunca miraba estos archivos, así que facturas cobradas volvían como vencidas mes tras mes.

La constancia nombra el comprobante exacto (`Serie de Comprobante` + `Numero de Comprobante`), así
que NO hay ambigüedad — a diferencia de casar por importe, que produce falsos (el abono de un cliente
calzaba con la factura de otro por el mismo monto).

Los archivos vienen de SOL → Consulta de detracciones, en .txt separado por barras; el .csv es la
misma data con comas. Se leen los DOS por si solo existe uno.
"""
from __future__ import annotations

import csv
import glob
import io
import os

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
PATRON = "0[12]. */03. DETRACCIONES/*/02. Constancias_SUNAT/DETR-SUNAT_*"


def depositadas(raiz: str | None = None) -> dict[str, dict]:
    """{comprobante: {monto, fecha, adquirente, constancia}} de toda detracción depositada.

    La clave es el comprobante normalizado (`F001-1803`), que es como lo guarda el maestro de
    cobros: la constancia trae la serie y el número por separado y el número con ceros a la
    izquierda (`00001803`).
    """
    base = raiz or RAIZ
    out: dict[str, dict] = {}
    for ruta in sorted(glob.glob(os.path.join(base, PATRON))):
        if not ruta.lower().endswith((".txt", ".csv")):
            continue
        deli = "|" if ruta.lower().endswith(".txt") else ","
        try:
            filas = list(csv.reader(io.open(ruta, encoding="latin-1"), delimiter=deli))
        except Exception:
            continue
        for r in filas:
            # la cabecera y las filas basura no tienen número de constancia numérico
            if len(r) < 16 or not r[2].strip().isdigit():
                continue
            try:
                num = int(r[15])
                monto = float(r[10])
            except (ValueError, IndexError):
                continue
            comp = f"{r[14].strip()}-{num}"
            # si el mismo comprobante aparece en dos meses, gana el depósito más reciente
            prev = out.get(comp)
            if prev is None or r[9] > prev["fecha"]:
                out[comp] = {"monto": monto, "fecha": r[9], "adquirente": r[8].strip(),
                             "constancia": r[2].strip()}
    return out


if __name__ == "__main__":
    d = depositadas()
    print(f"{len(d)} comprobantes con detracción depositada")
    for k, v in sorted(d.items())[:5]:
        print(f"  {k:12s} S/{v['monto']:>8,.2f}  {v['fecha']}  {v['adquirente'][:30]}")


def cobradas_bruto(raiz: str | None = None) -> dict[str, str]:
    """{comprobante: fecha de cobro} de las facturas que el detractor marcó con cobro BRUTO.

    Segunda prueba de cobro, complementaria a la constancia. Una auto-detracción SOLO nace porque
    el cliente pagó el importe completo sin retener el 12%: si la hoja 09 la registra, la factura
    se cobró — aunque todavía no exista constancia, porque el depósito lo debemos nosotros.

    Sin esto quedaban abiertas facturas ya cobradas que el sistema insistía en sacar
    como vencidas. El caso es justo el inverso al de
    `depositadas()`: allí el 12% ya está en el BN; aquí lo debemos y por eso no hay constancia.
    """
    import sys
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from lib import sources as S
    import openpyxl

    try:
        wb = openpyxl.load_workbook(S.latest("detracciones"), data_only=True)
    except Exception:
        return {}
    hojas = [n for n in wb.sheetnames if n.startswith("09")]
    if not hojas:
        return {}
    import re as _re
    out: dict[str, str] = {}
    for r in wb[hojas[0]].iter_rows(values_only=True):
        if not r:
            continue
        # Solo el bloque A (auto-detracciones nuestras). Leer la hoja entera hacía que la frase
        # «los TRES pagaron BRUTO» de una fila del bloque F marcara una factura que pagó
        # NETO. Los bloques B–H empiezan con "B) ", "C) "…
        cab = str(r[0] or "").strip()
        if _re.match(r"^[B-H]\)", cab):
            break
        if len(r) < 11:
            continue
        comp = str(r[4] or "")
        obs = " ".join(str(c) for c in r[7:] if c).upper()
        if (comp.startswith(("F001-", "E001-")) and _re.search(r"\bBRUTO\b", obs)
                and "NETO" not in obs and "ANULADA" not in obs):
            out[comp] = str(r[2] or "")[:10]
    return out
