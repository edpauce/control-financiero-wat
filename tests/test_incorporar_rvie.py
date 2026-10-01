"""
incorporar_rvie (ventas RVIE → COBROS): un abono del banco paga UNA factura.

  · casar_cobro marca el abono como usado: la gemela que viene detrás no lo reclama;
  · un abono que ya estaba reclamado (usado=True) no se reasigna;
  · una NC sobre una factura YA COBRADA no toca el cobro: se marca como devolución a validar;
  · lo que ya está en el maestro se salta (dedupe), no se reescribe su cobro.
Solo tests: el módulo no se toca. Todo con libros sintéticos en tmp_path.
"""
import importlib.util
import os
from datetime import date

import openpyxl

_ruta = os.path.join(os.path.dirname(__file__), "..", "tools", "incorporar_rvie.py")
_spec = importlib.util.spec_from_file_location("incorporar_rvie", _ruta)
RV = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(RV)


def _doc(comp="F001-1045", total=1000.0, femi=date(2026, 7, 1), **kw):
    d = {"periodo": "2026-07", "comp": comp, "tipo": "01", "femi": femi, "ruc": "20100000001",
         "cli": "JURISBA EIRL", "bi": total / 1.18, "igv": total - total / 1.18, "total": total,
         "moneda": "PEN", "es_nc": total < 0, "modifica": ""}
    d.update(kw)
    return d


def _abono(monto, f=date(2026, 7, 10), usado=False):
    return {"f": f, "m": monto, "desc": "TRANSF JURISBA", "ref": "", "usado": usado}


# ─── casar_cobro ─────────────────────────────────────────────────────────────
def test_un_abono_paga_una_sola_factura():
    """Dos facturas gemelas (mismo importe) y un solo abono neto: la segunda queda por cobrar."""
    abonos = [_abono(880.0)]
    a, modo = RV.casar_cobro(_doc("F001-1045"), abonos)
    assert (a["m"], modo, a["usado"]) == (880.0, "neto", True)
    assert RV.casar_cobro(_doc("F001-1046"), abonos) is None


def test_abono_ya_reclamado_no_se_reasigna():
    assert RV.casar_cobro(_doc(), [_abono(880.0, usado=True)]) is None
    assert RV.casar_cobro(_doc(), [_abono(1000.0, usado=True)]) is None


def test_neto_manda_sobre_bruto_y_la_ventana_es_menos5_mas120_dias():
    a, modo = RV.casar_cobro(_doc(), [_abono(1000.0), _abono(880.0)])
    assert (a["m"], modo) == (880.0, "neto")
    assert RV.casar_cobro(_doc(), [_abono(1000.0)])[1] == "bruto"
    assert RV.casar_cobro(_doc(), [_abono(880.0, f=date(2026, 6, 25))]) is None      # 6 días antes
    assert RV.casar_cobro(_doc(), [_abono(880.0, f=date(2026, 6, 27))]) is not None  # 4 días antes
    assert RV.casar_cobro(_doc(), [_abono(880.0, f=date(2026, 10, 30))]) is None     # 121 días después


def test_nc_dolares_o_sin_fecha_no_casan():
    assert RV.casar_cobro(_doc(total=-1000.0), [_abono(-1000.0)]) is None
    assert RV.casar_cobro(_doc(moneda="USD"), [_abono(880.0)]) is None
    assert RV.casar_cobro(_doc(femi=None), [_abono(880.0)]) is None


def test_bajo_el_umbral_spot_no_hay_detraccion_que_restar():
    a, modo = RV.casar_cobro(_doc(total=700.0), [_abono(700.0)])
    assert (a["m"], modo) == (700.0, "neto")       # total ≤ 700 → neto == bruto


# ─── incorporar: el maestro no pierde cobros confirmados ─────────────────────
def _cobros(ruta):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("MERAK - Cobros")
    for _ in range(3):
        ws.append([None] * RV.NCOLS)
    ws.append(["Jul", date(2026, 7, 1), "F001-1045", "JURISBA EIRL", 1000.0, 0.12, 120.0, 880.0,
               "CREDITO", date(2026, 7, 10), 1000.0, 0.0, "Cobrado", None, "cobro bruto confirmado"])
    ws.append(["TOTALES"] + [None] * (RV.NCOLS - 1))
    wb.create_sheet("ALTAMIRA - Cobros").append([None] * RV.NCOLS)
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    wb.save(ruta)
    return ruta


def _preparar(tmp_path, monkeypatch):
    base = _cobros(os.path.join(str(tmp_path), "04. FINANZAS", "08. COBROS", "COBROS_2026-07.xlsx"))
    monkeypatch.setattr(RV, "ROOT", str(tmp_path))
    monkeypatch.setattr(RV.S, "latest", lambda clave: base if clave == "cobros" else None)


def _fila(ws, comp):
    return next(r for r in range(4, ws.max_row + 1) if RV._norm(ws.cell(r, 3).value) == comp)


def test_nc_sobre_factura_cobrada_no_toca_el_cobro(tmp_path, monkeypatch):
    _preparar(tmp_path, monkeypatch)
    nc = _doc("E001-160", total=-1000.0, tipo="07", femi=date(2026, 7, 20), modifica="F001-1045")
    r = RV.incorporar("MER", [nc], [], aplicar=True)
    ws = openpyxl.load_workbook(r["out"])["MERAK - Cobros"]
    f = _fila(ws, "F001-1045")
    assert (ws.cell(f, 11).value, ws.cell(f, 12).value, ws.cell(f, 13).value) == (1000.0, 0.0, "Cobrado")
    assert str(ws.cell(f, 15).value).startswith("⚠️ NC E001-160") and "DEVOLUCIÓN" in str(ws.cell(f, 15).value)
    assert r["nc_ok"] == [] and r["nc_huerf"] == [("E001-160", "anula F001-1045 que ya estaba cobrada — devolución")]


def test_factura_ya_en_el_maestro_se_salta_y_conserva_su_cobro(tmp_path, monkeypatch):
    _preparar(tmp_path, monkeypatch)
    r = RV.incorporar("MER", [_doc("F001-1045")], [_abono(880.0)], aplicar=True)
    assert (r["saltados"], r["nuevos"]) == (1, 0)
    ws = openpyxl.load_workbook(r["out"])["MERAK - Cobros"]
    f = _fila(ws, "F001-1045")
    assert (ws.cell(f, 10).value.date(), ws.cell(f, 11).value, ws.cell(f, 13).value) == \
        (date(2026, 7, 10), 1000.0, "Cobrado")


def test_nc_sin_destino_no_se_aplica_por_monto(tmp_path, monkeypatch):
    """Con gemelas, adivinar por importe anula la equivocada."""
    _preparar(tmp_path, monkeypatch)
    nc = _doc("E001-160", total=-1000.0, tipo="07", femi=date(2026, 7, 20), modifica="")
    r = RV.incorporar("MER", [nc], [], aplicar=True)
    ws = openpyxl.load_workbook(r["out"])["MERAK - Cobros"]
    assert ws.cell(_fila(ws, "F001-1045"), 13).value == "Cobrado"
    assert r["nc_huerf"] == [("E001-160", "sin 'Nro CP Modificado' en el RVIE")]


def test_venta_nueva_casa_con_el_banco_y_marca_el_abono(tmp_path, monkeypatch):
    _preparar(tmp_path, monkeypatch)
    abonos = [_abono(1760.0, f=date(2026, 7, 15))]
    r = RV.incorporar("MER", [_doc("F001-1050", total=2000.0, femi=date(2026, 7, 5))], abonos, aplicar=True)
    assert (r["nuevos"], r["cobrado_n"]) == (1, 1) and abonos[0]["usado"] is True
    ws = openpyxl.load_workbook(r["out"])["MERAK - Cobros"]
    f = _fila(ws, "F001-1050")
    assert [ws.cell(f, c).value for c in (5, 7, 8, 11, 12, 13)] == [2000.0, 240.0, 1760.0, 1760.0, 0.0, "Cobrado"]


def test_sin_aplicar_no_escribe_nada(tmp_path, monkeypatch):
    """El gate exige aprobación para movimientos grandes. Si la simulación también la exigiera,
    la única forma de MIRAR una corrida de S/60,000 sería autorizarla — que es lo contrario de
    lo que el umbral existe para lograr."""
    _preparar(tmp_path, monkeypatch)
    r = RV.incorporar("MER", [_doc("F001-9999", total=30000.0)], [], aplicar=False)
    assert r.get("preview") is True
    assert "out" not in r, "la simulación devolvió una ruta escrita"
    assert r["nuevos_n"] == 1 and r["nuevos_total"] == 30000.0
