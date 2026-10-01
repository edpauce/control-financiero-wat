"""
Mejora 4 (10/08) — Sello de integridad: los NÚMEROS del banco de un mes sellado no cambian
sin gritar; la capa de clasificación (categoría/referencia/obs) queda libre — es la que se
enriquece en sitio. Nace del incidente del 10/08 (una regeneración pisó julio formateado).
"""
import importlib.util
import os

import openpyxl

_ruta = os.path.join(os.path.dirname(__file__), "..", "tools", "lib", "sellos.py")
_spec = importlib.util.spec_from_file_location("sellos", _ruta)
SE = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(SE)


def _wb():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ver = wb.create_sheet(SE.HOJAS_MOV["MERAK"])
    for _ in range(4):
        ver.append([None] * 9)
    ver.append(["02/07/2026", "Julio", "DE CLIENTE X", "INGRESO CLIENTE",
                4100.0, None, 104.0, "CLIENTE X", "op 1"])
    ver.append(["15/07/2026", "Julio", "A 193 1 1", "PAGO PROVEEDOR / TERCERO",
                None, 1000.0, -896.0, "PROVEEDOR Y", "op 2"])
    ver.append(["13/06/2026", "Junio", "DE OTRO", "INGRESO CLIENTE",
                500.0, None, 500.0, "OTRO", "op 3"])
    ver.append(["TOTALES JULIO 2026 (al 31/07)", None, None, None, None, None, None, None, None])
    wb.create_sheet(SE.HOJAS_MOV["ALTAMIRA"]).append([None] * 9)
    return wb


def test_sellar_y_verificar_limpio():
    wb = _wb()
    sellados = SE.sellar(wb, cuando="2026-08-10")
    assert ("MERAK", "Julio") in [(e, m) for e, m, _ in sellados]
    assert ("MERAK", "Junio") in [(e, m) for e, m, _ in sellados]
    assert SE.verificar(wb) == []
    assert wb[SE.HOJA_SELLOS].sheet_state == "hidden"


def test_alterar_un_numero_grita():
    wb = _wb()
    SE.sellar(wb)
    ws = wb[SE.HOJAS_MOV["MERAK"]]
    ws.cell(6, 6).value = 999.0                      # cargo 1,000 → 999
    v = SE.verificar(wb)
    assert len(v) == 1 and "MERAK Julio" in v[0] and "cargos" in v[0]


def test_la_capa_de_clasificacion_es_libre():
    """Categoría, referencia y observación se enriquecen en sitio sin romper el sello."""
    wb = _wb()
    SE.sellar(wb)
    ws = wb[SE.HOJAS_MOV["MERAK"]]
    ws.cell(5, 4).value = "OTRA CATEGORÍA"
    ws.cell(5, 8).value = "NOMBRE NUEVO (cola de identificación)"
    ws.cell(5, 9).value = "op 1 · paga F001-1074 (cruce)"
    assert SE.verificar(wb) == []


def test_borrar_el_mes_sellado_grita_y_resellar_acepta():
    wb = _wb()
    SE.sellar(wb)
    ws = wb[SE.HOJAS_MOV["MERAK"]]
    ws.cell(7, 2).value = "X"                        # la fila de junio deja de ser 'Junio'
    v = SE.verificar(wb)
    assert any("Junio" in x and "YA NO TIENE FILAS" in x for x in v)
    SE.sellar(wb)                                    # re-sellar = aceptar conscientemente
    assert all("Junio" not in x for x in SE.verificar(wb))


def test_totales_y_clasificacion_no_entran_a_la_huella():
    h1 = SE.huella(_wb(), "MERAK", "Julio")
    wb = _wb()
    ws = wb[SE.HOJAS_MOV["MERAK"]]
    ws.cell(8, 1).value = "TOTALES JULIO 2026 (otro texto)"   # banner distinto
    ws.cell(5, 4).value = "RECLASIFICADA"                     # clasificación distinta
    assert SE.huella(wb, "MERAK", "Julio")["sha"] == h1["sha"]
