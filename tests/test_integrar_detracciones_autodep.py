"""
integrar_detracciones: el ROSADO de la hoja BN lo decide AUTODEP_DOCS de la config del mes.

Rosado = auto-depósito de la empresa (es plata nuestra que entró al BN, no una detracción del
cliente). Si el tool pintara por su cuenta, o dejara de pintar lo que la config dice, el
compliance SPOT del mes se lee mal. Todo en tmp_path: no toca los maestros reales.
"""
import importlib.util
import os

import openpyxl
import pytest

from lib import cierre as CI

ROSADO = "F4CCCC"


def _tool(nombre):
    ruta = os.path.join(os.path.dirname(__file__), "..", "tools", f"{nombre}.py")
    spec = importlib.util.spec_from_file_location(nombre, ruta)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _base(ruta):
    """SEGUIMIENTO mínimo: hojas BN con una banda de junio (para copiar estilos) y la hoja 09."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for hoja in ("04 MERAK BN", "05 ALTAMIRA BN"):
        ws = wb.create_sheet(hoja)
        ws.append(["TOTAL Ene-Jun", None, None, None, 0, 0, 0])
        ws.append(["Junio 2026 (al 30/06)", None, None, "— Movimientos BN junio"])
        for _ in range(2):
            ws.append([None] * 7)
        ws.append(["Jun 2026", "05/06/2026", "DOC-VIEJO/LIMA", "DEPOSITO", None, 10.0, 100.0])
        ws.append(["Subtotal Jun2026", "1 movimientos", None, None, 0, 10.0, 100.0])
    ws9 = wb.create_sheet("09 Pendientes Auto-Det")
    ws9.append(["Pendientes"])
    ws9.append(["nota"])
    ws9.append(["Saldo BN MERAK al 30/06/2026 (S/)", None, None, None, 0])
    ws9.append(["Saldo BN ALTAMIRA al 30/06/2026 (S/)", None, None, None, 0])
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    wb.save(ruta)
    return ruta


REP_BN = {"movimientos": [
    {"nro": 2, "fecha": "2026.07.20", "trans": "DEPOSITO", "documento": "DOC-CLI", "ruc": "",
     "oficina": "LIMA", "cargo": 0.0, "abono": 240.0},
    {"nro": 1, "fecha": "2026.07.05", "trans": "DEPOSITO", "documento": "DOC-AUTO", "ruc": "",
     "oficina": "CUSCO", "cargo": 0.0, "abono": 120.0},
], "n": 2, "abonos": 360.0, "cargos": 0.0}


def _correr(tmp_path, monkeypatch, cfg_extra=""):
    ID = _tool("integrar_detracciones")
    monkeypatch.setattr(CI, "CONFIG_DIR", str(tmp_path))
    (tmp_path / "2026-07.py").write_text(
        'SALDO_BN = {"MERAK": 40000.0, "ALTAMIRA": 1500.0}\nREV_FECHA = "2026-07-31"\n' + cfg_extra)
    monkeypatch.setattr("sys.argv", ["integrar_detracciones.py", "--mes", "2026-07"])
    monkeypatch.setattr(ID, "ROOT", str(tmp_path))
    d = os.path.join(str(tmp_path), "04. FINANZAS", "10. DETRACCIONES")
    base = _base(os.path.join(d, "BASE_SIN_JULIO.xlsx"))
    monkeypatch.setattr(ID, "base_seguimiento", lambda cfg: base)
    monkeypatch.setattr(ID.CI, "first", lambda pat: "bn_julio.xlsx")
    monkeypatch.setattr(ID.BN, "parse_movimientos_bn", lambda ruta: REP_BN)
    out = ID.main()
    return openpyxl.load_workbook(out)


def _fila_doc(ws, doc):
    return next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 3).value or "").startswith(doc))


def _fill(ws, r):
    f = ws.cell(r, 1).fill
    return str(f.start_color.rgb)[-6:] if f is not None and f.fill_type else ""


def test_autodep_docs_manual_pinta_solo_ese_documento(tmp_path, monkeypatch):
    wb = _correr(tmp_path, monkeypatch, 'AUTODEP_DOCS = {"DOC-AUTO"}\n')
    for hoja in ("04 MERAK BN", "05 ALTAMIRA BN"):
        ws = wb[hoja]
        assert _fill(ws, _fila_doc(ws, "DOC-AUTO")) == ROSADO
        assert _fill(ws, _fila_doc(ws, "DOC-CLI")) != ROSADO
        assert ws.cell(_fila_doc(ws, "DOC-AUTO"), 6).value == 120.0


def test_sin_autodep_docs_nada_es_rosado(tmp_path, monkeypatch):
    wb = _correr(tmp_path, monkeypatch)
    ws = wb["04 MERAK BN"]
    assert _fill(ws, _fila_doc(ws, "DOC-AUTO")) != ROSADO
    assert _fill(ws, _fila_doc(ws, "DOC-CLI")) != ROSADO


def test_banda_del_mes_subtotal_y_saldo_bn(tmp_path, monkeypatch):
    wb = _correr(tmp_path, monkeypatch, 'AUTODEP_DOCS = {"DOC-AUTO"}\n')
    ws = wb["04 MERAK BN"]
    banda = next(r for r in range(1, ws.max_row + 1)
                 if str(ws.cell(r, 1).value or "").startswith("Julio 2026"))
    # cronológico ascendente: el 05/07 antes que el 20/07 (el parser los trae al revés)
    assert ws.cell(banda + 1, 2).value == "05/07/2026" and ws.cell(banda + 2, 2).value == "20/07/2026"
    sub = ws.cell(banda + 3, 1).value
    assert sub == "Subtotal Jul2026" and ws.cell(banda + 3, 6).value == 360.0
    assert ws.cell(banda + 3, 7).value == 40000.0, "el saldo BN del subtotal es el de la config"
    ws9 = wb["09 Pendientes Auto-Det"]
    assert ws9.cell(3, 5).value == 40000.0 and ws9.cell(4, 5).value == 1500.0
    assert "31/07/2026" in str(ws9.cell(3, 1).value)


def test_falta_saldo_bn_aborta(tmp_path, monkeypatch):
    ID = _tool("integrar_detracciones")
    monkeypatch.setattr(CI, "CONFIG_DIR", str(tmp_path))
    (tmp_path / "2026-07.py").write_text("")
    monkeypatch.setattr("sys.argv", ["integrar_detracciones.py", "--mes", "2026-07"])
    with pytest.raises(SystemExit) as e:
        ID.main()
    assert "SALDO_BN" in str(e.value)
